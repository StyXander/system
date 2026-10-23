"""重新打开真实导出，并核对字体及CSV/JSON关系同源性。"""
import csv,json
from pathlib import Path
from docx import Document
ROOT=Path(__file__).resolve().parents[1];folder=ROOT/'outputs/financial-impact-acceptance/browser';csv.field_size_limit(10000000)
run=json.loads((folder/'export.json').read_text(encoding='utf-8'))
rows=list(csv.reader((folder/'export.csv').open(encoding='utf-8-sig',newline='')))
actual=[json.loads(row[11]) for row in rows if row and row[0]=='financial_relation'];expected=run['financial_impact']['relations']
doc=Document(folder/'export.docx');bad=[];texts=[]
for paragraph in [*doc.paragraphs,*[q for t in doc.tables for row in t.rows for c in row.cells for q in c.paragraphs]]:
 heading=(paragraph._p.style or '').startswith('Heading') or paragraph._p.style=='Title'
 for part in paragraph.runs:
  if part.text.strip() and (part.font.name!=('黑体' if heading else '宋体') or part.font.size.pt!=(16 if heading else 12)):bad.append(part.text[:20])
 texts.append(paragraph.text)
receipt={'docx_opened':True,'paragraphs':len(texts),'tables':len(doc.tables),'font_violations':len(bad),'has_financial_section':any('事项、会计路径、附注项目' in t for t in texts),'financial_csv_relations':len(actual),'financial_csv_json_identical':actual==expected,'financial_item_source_edges_correct':all(ref['document_id']==next(item['document_id'] for item in run['financial_impact']['items'] if item['item_id']==relation['financial_item_id']) and ref['pdf_page']==next(item['pdf_page'] for item in run['financial_impact']['items'] if item['item_id']==relation['financial_item_id']) for relation in expected if relation.get('financial_item_id') for ref in relation['supporting']),'notice':'结构、字体与同源性检查，不冒充全页排版审定'}
(folder/'export-verification.json').write_text(json.dumps(receipt,ensure_ascii=False,indent=2),encoding='utf-8');print(receipt)
assert not bad and actual==expected and receipt['financial_item_source_edges_correct']
