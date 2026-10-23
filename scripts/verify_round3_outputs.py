"""第三轮交付核验：六级分档、AI 判定与数字/引文的导出一致性。

本脚本是「口径继承与证据分级」方案（来源一）在第三轮登记的核验器。
原文件被 `.gitignore` 的 `/scripts/*` 整体忽略而未进入提交，导致
`backend/tests/test_round3_export_parity.py` 在收集阶段即报
FileNotFoundError。此处依据该测试的调用契约、`connectivity_overview`
的六级定义与 `connectivity_delivery` 的真实导出接口重建。

重建原则：
- 只校验可回查事实，不修改运行结果、不改写历史 JSON；
- 分档取自 `build_overview` 的 `program_levels`，不在本地另立一套口径；
- 数字核对区分「编造」与「单位换算」，避免把 100 万元→1000000 元 误判为编造。
"""
from __future__ import annotations

import re
from decimal import Decimal, InvalidOperation
from io import BytesIO

from docx import Document

from app.connectivity_delivery import AI_LABELS
from app.connectivity_delivery import export_run as _delivery_export_run
from app.connectivity_overview import build_overview

# 六级互斥分类，键必须与 connectivity_overview.build_overview['program_levels'] 一致。
LEVEL_LABELS = {
    "formal": "正式一致",
    "conditional": "假设内相容（条件试算）",
    "limited": "有限支持",
    "insufficient": "依据不足",
    "conflict": "已确认冲突或未解释差异",
    "not_applicable": "不适用数值计算",
}

EXPORT_FORMATS = ("json", "csv", "docx")

_UNIT_FACTOR = {
    "元": Decimal("1"), "千元": Decimal("1000"), "万元": Decimal("10000"),
    "百万元": Decimal("1000000"), "亿元": Decimal("100000000"),
    "吨": Decimal("1"), "万吨": Decimal("10000"),
}


def export_run(run, format_name):
    """导出入口。测试通过替换本符号制造「标签塌缩」以验证核验器本身有效。"""
    return _delivery_export_run(run, format_name)


def _compact(text: str) -> str:
    return re.sub(r"\s+", "", str(text or ""))


def _docx_text(payload: bytes) -> str:
    document = Document(BytesIO(payload))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def _payload_text(payload: bytes, format_name: str) -> str:
    if format_name == "docx":
        return _docx_text(payload)
    return payload.decode("utf-8", errors="replace")


def _levels(run) -> dict:
    overview = run.overview or build_overview(run)
    program = overview["program_levels"]
    return {key: program[key]["value"] for key in LEVEL_LABELS}


def inspect_run(run) -> dict:
    """检视运行本身的完整性：关系身份、引文可回溯、数字有出处。"""
    problems: list[str] = []

    seen: dict[str, int] = {}
    for check in run.checks:
        seen[check.check_id] = seen.get(check.check_id, 0) + 1
    for check_id, count in seen.items():
        if count > 1:
            problems.append(f"重复 check_id：{check_id} 出现 {count} 次；重复关系会虚增分档计数")

    occurrences = {item.occurrence_id: item for item in run.occurrences}
    for check in run.checks:
        for condition in check.conditions:
            for ref in condition.evidence:
                target = occurrences.get(ref.occurrence_id)
                if target is None:
                    problems.append(f"引用不能逐字回溯：{check.check_id}/{condition.name} 指向不存在的披露 {ref.occurrence_id}")
                elif _compact(ref.quote) not in _compact(target.context):
                    problems.append(f"引用不能逐字回溯：{check.check_id}/{condition.name} 的引文不在原文上下文中")
        for ref in check.relationship_evidence:
            target = occurrences.get(ref.occurrence_id)
            if target is None or _compact(ref.quote) not in _compact(target.context):
                problems.append(f"引用不能逐字回溯：{check.check_id} 的关系依据无法回到原文")

    for opinion in run.ai_analyses:
        for ref in opinion.supporting_evidence:
            target = occurrences.get(ref.occurrence_id)
            if target is None:
                problems.append(f"引用不能逐字回溯：{opinion.check_id} 的 AI 依据指向不存在的披露 {ref.occurrence_id}")
            elif _compact(ref.quote) not in _compact(target.context):
                problems.append(f"引用不能逐字回溯：{opinion.check_id} 的 AI 引文不在原文上下文中")

    allowed = _bound_numbers(run)
    for opinion in run.ai_analyses:
        for token in _numbers(opinion.summary):
            if _normalize(token) not in allowed:
                problems.append(f"未绑定数值：{opinion.check_id} 的 AI 说明出现原文与可复算换算都没有的数字 {token}")

    return {"problems": problems, "levels": _levels(run), "ai_count": len(run.ai_analyses)}


def _numbers(text: str) -> list[str]:
    return re.findall(r"\d+(?:\.\d+)?", str(text or ""))


def _normalize(token: str) -> str:
    try:
        value = Decimal(token.replace(",", ""))
    except (InvalidOperation, AttributeError):
        return str(token)
    if value == value.to_integral_value():
        return str(value.quantize(Decimal(1)))
    return format(value.normalize(), "f")


def _bound_numbers(run) -> set[str]:
    """原文数字 + 由原值按已声明单位换算后的等值，二者都算有出处。"""
    allowed: set[str] = set()
    for item in run.occurrences:
        if item.value:
            for token in _numbers(item.value):
                allowed.add(_normalize(token))
            try:
                base = Decimal(str(item.value).replace(",", ""))
            except (InvalidOperation, ValueError):
                base = None
            factor = _UNIT_FACTOR.get(item.unit or "")
            if base is not None and factor:
                allowed.add(_normalize(str(base * factor)))
        for token in _numbers(item.context):
            allowed.add(_normalize(token))
    return allowed


def verify_exports(run) -> dict:
    """确认六级分档不合并、AI 判定不塌缩、非法控制字符不改写结果。"""
    problems: list[str] = []
    before = run.model_dump(mode="json")
    levels = _levels(run)
    ai_count = len(run.ai_analyses)

    for format_name in EXPORT_FORMATS:
        try:
            payload, _mime = export_run(run, format_name)
        except Exception as error:  # 导出崩溃本身就是交付缺陷
            problems.append(f"{format_name} 导出失败：{type(error).__name__}: {error}")
            continue
        text = _payload_text(payload, format_name)
        # JSON 保留模型判定的原始枚举值，CSV/Word 才呈现中文标签；两者都必须可区分。
        for verdict, label in AI_LABELS.items():
            expected = sum(1 for item in run.ai_analyses if item.verdict == verdict)
            if not expected:
                continue
            needle = f'"verdict":"{verdict}"' if format_name == "json" else label
            haystack = _compact(text) if format_name == "json" else text
            if haystack.count(needle) < expected:
                problems.append(f"{format_name} 的 AI 判定被合并：{label} 应出现 {expected} 次，实际 {haystack.count(needle)} 次")

    after = run.model_dump(mode="json")
    if before != after:
        problems.append("导出过程改写了运行数据；导出必须只读")

    overview = run.overview or build_overview(run)
    if overview["program_levels"] and sum(levels.values()) != len(run.checks):
        problems.append(f"六级分档合计 {sum(levels.values())} 与关系总数 {len(run.checks)} 不一致")

    return {"passed": not problems, "problems": problems, "levels": levels, "ai_count": ai_count}
