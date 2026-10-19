"""在原模板上生成修订候选，保留源文件；Office渲染未通过前不称正式提交件。"""
from __future__ import annotations

import argparse
import json
from pathlib import Path
import zipfile
import xml.etree.ElementTree as ET

from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
NOTICE = 'AI生成内容，用于年度报告与可持续发展报告关联核对和后续调查参考，不构成审计结论、鉴证结论或鉴证意见。'
SCOPE = '基于已登记、公开的年度报告与可持续发展报告，辅助专业人员判断披露关系是否可比、完成可复算检查、定位证据及待补条件。'

REPLACEMENTS = {
 '面向审计计划阶段的': '年度报告与可持续披露的',
 '证据约束型多智能体辅助系统': '专业关联核对辅助系统',
 '竞赛答辩  /  销售与收款循环': '竞赛答辩  /  公开双报告核验',
 '审迹智链 · 审计计划阶段': '审迹智链 · 公开双报告核验',
 '销售与收款循环': '年报与可持续披露',
 '异常线索': '披露关系',
 '为审计人员能够继续核查的材料。': '为专业人员能够继续核对的记录。',
 '会计师事务所审计项目组': '公开报告核对人员',
 '，在计划阶段明确重点核查方向，业务承接与续聘暂列未来路线。': '，先判断数字能否比较，再解释差异及所需证据。',
 '审计计划阶段的用户任务': '两份报告之间的核对任务',
 '审计人员需要把': '使用者需要把',
 '业务了解和报表变化': '原文数字和比较条件',
 '核查安排': '关联记录',
 '围绕收入与应收账款变化寻找待核查线索，结合披露信息理解异常，不能仅凭规则触发认定舞弊。': '检查同值、组成与比例关系，保留同名但不同范围的反例，不能将不同数字直接认定为企业错误。',
 '记录依据、资料缺口和后续核查动作，由项目组决定核查优先级，并在取得补充证据后调整计划。': '逐项记录原文、公式、差额和缺少的定义或范围说明；补证后重新核对，保留此前版本。',
 '应收账款增速高于收入只是背离，不能直接断定收入虚增；须引入相反解释与资料缺口，明确待核查事项，结论由注册会计师作出。': '集团与母公司范围不同、金额与人数不同都不能直接比较；模型意见不能覆盖程序门禁，也不能替代专业判断。',
 '把待核查事项对应到认定、需要补充的资料及建议程序，帮助项目组决定下一步查什么、向谁索取。': '把待查关系对应到缺少的定义、统计范围或组成说明，解释补到哪些材料才能改变当前判断。',
 '以确定性规则定位线索，以证据限制解释范围，再由审计人员回查原文并决定后续安排。': '以登记原件形成关系候选，以口径条件限制比较，再由程序计算和专业人员回查原文。',
 '规则预筛': '关系计算',
 'R1比较收入与应收账款增长变化；辅助规则结合可比性处理。计算保留输入与口径，规则不认定舞弊。': '六类登记规则处理同值、组成、比例、多位置与叙述关系。计算保留输入、单位和前提，缺口径时只作条件试算。',
 '按案例检索披露和相关规范，围绕异常组织质疑、反证与复核，未获支持的解释明确标注为待验证假设。': '按本次事项检索报告原文，组织候选、反证和复核。历史知识库未接入当前主链，不能用其规模宣传新能力。',
 '输出预审风险备忘录、Top 5待核查事项和待索取资料清单，保留引用、资料缺口及建议核查程序。': '输出关联核对记录、待调查事项和具体补证动作，保留原页、程序事实、AI意见及复核记录。',
 '系统工作台': '历史业务界面',
 '已有真实界面截图': '不作为本轮双报告验收截图',
 '历史运行示例': '历史业务示例，非本轮结果',
 '登记案例和来源快照，载入字段及原文位置。': '读取已登记双报告，检查原件哈希和来源页。',
 '核对口径并计算指标，形成可解释的异常线索。': '检查主体、期间、定义、范围、单位与关系。',
 '在当前案例范围内检索，记录命中来源与缺口。': '围绕当前关系补查报告证据，保留未找到项。',
 '质疑提出问题，反证比较解释，复核检查依据。': '候选、定向反证、复核分阶段组织模型意见。',
 '输出核查事项与资料清单，附运行记录供回查。': '同源输出关系、待查事项与具体补证动作。',
 '结合规则结果提出具体核查问题，说明涉及的认定及支持材料；证据不足时保留疑点，不作最终定性。': '结合报告原文提出关系候选及比较条件；金额、人数、期间和成员必须通过程序验证，失败项单独记录。',
 'FastAPI组织六阶段任务，记录角色尝试和校验结果；模型服务采用DeepSeek官方直连，调用状态单独留痕。': 'FastAPI组织任务和阶段记录；模型调用受IP、并发与每日预算约束。当前本地固定任务没有模型调用。',
 '确定性规则、案例隔离检索、三角色协作与硬校验按顺序衔接，数字输入和解释主张使用不同的支持要求。': '表格抽取、关系候选、比较条件和确定性计算相互衔接；数字事实与AI意见各自保留状态及证据。',
 '来源快照、哈希、引用标识与运行台账支撑复验；公开演示完成结果写入Supabase，运行中断需显式重跑。': '登记来源、哈希和运行台账支撑复验；本轮新固定任务保存在本地，线上Supabase与部署状态另行核验。',
 '四项围绕证据的创新控制': '可演示机制与待验证增量',
 '创新重点落在主张如何获得支持，以及证据不足时系统怎样收缩解释范围。': '组合关注指标可比性、分级复算、反证解释和补证动作。增益需要同样本对照，不宣称行业首创。',
 '规则驱动路由': '比较条件门禁',
 '由确定性结果选择需要解释的异常，减少脱离数据的泛泛分析，使角色任务与具体线索对应。': '先核查主体、期间、定义、范围和单位；前提缺失时不把算术相近提升为正式相符。',
 '认定—证据—程序': '公式与条件试算',
 '用覆盖矩阵连接涉及的认定、已有证据和建议核查程序，显式保留资料缺口与下一步动作。': '程序保存原值、单位、成员及公式；正式残差和有前提的试算分开，显示精度不等于已证实舍入。',
 '证据适配度': '解释与反证',
 '区分原始字段、计算证据与解释材料，限制每类证据能够支持的主张，避免一条引用承担过多判断。': '解释需绑定具体原文并检查竞争假设。引用存在不等于推理正确，反证撤回能力仍需专项验证。',
 '数字与反证回查': '调查动作与复现',
 '将数字主张回查到输入或原文，同时记录相反解释的检索轨迹，降低只搜支持证据的倾向。': '待查项应说明缺少哪条定义或附注；JSON、Word和本地任务来自同一保存结果，保持不同运行性质。',
 'B0为人工基线，B1仅做规则计算，B2加入单模型草稿，B3采用完整协作链；不同组的失败情况也纳入记录。': '关联实验另设A纯规则、B单模型、C旧三阶段、D升级三阶段；历史B0至B3定义不改。本轮只执行零调用程序组。',
 '当前发布记录登记23个演示案例，原型包含规则、检索、角色协作与报告链；该数量表示案例入口规模。': '当前登记5对10份报告、4家公司、3主题、6类规则。数量由服务端登记事实生成，不等于全指标覆盖或专业验证。',
 '补齐人工基线与正式评分，核验生产模型完整执行及异常恢复；承接、续聘和更广场景适配保留为未来方向。': '三个业务分支已合入，工程与线上验收分别记录；双人标注、真实用户测试和付费模型对照尚未完成。',
 '依据当前发布记录：整改候选，人工评分待完成。': 'AI4候选稿：未部署，人工评分与Office渲染待完成。',
 '让审计计划阶段的线索有据可查': '让两份报告之间的关系有据可查',
 '让后续核查有明确的起点': '保留比较条件与具体补证动作',
}


def presentation(source, target):
    with zipfile.ZipFile(source) as package, zipfile.ZipFile(target, 'w', zipfile.ZIP_DEFLATED) as output:
        for info in package.infolist():
            body = package.read(info.filename)
            if info.filename.startswith('ppt/slides/slide') and info.filename.endswith('.xml'):
                tree = ET.fromstring(body)
                texts = [t for t in tree.iter() if t.tag.endswith('}t')]
                for element in texts:
                    text = element.text or ''
                    element.text = REPLACEMENTS.get(text, text).replace('/ 14', '/ 15')
                if info.filename == 'ppt/slides/slide14.xml':
                    # 学校与个人贡献页改成验证边界页，不制造人员履历或虚构人工贡献。
                    for element in texts:
                        element.text = ''
                    copy = {0:'14', 2:'验证边界与后续工作', 3:'VALIDATION BOUNDARIES', 4:'审迹智链 · 公开双报告核验',
                            5:'14', 7:' / 15', 8:'当前状态', 9:'程序检查、真实调用与人工评价分别记录。',
                            10:'本轮不以历史截图或零差额替代业务效果。', 12:'程序基线', 13:'17个冻结样例',
                            14:'13项机器预期符合', 15:'4项跨域验证待完成；专业标签待两人独立标注。',
                            18:'真实报告运行', 19:'五粮液2025', 20:'无模型计算', 21:'新本地任务可回读和导出；正式复算仍为0，不能称企业无异常。',
                            26:'模型与真人评价', 27:'待实测', 28:'不填写准确率或节时比例', 29:'单模型和三阶段对照需冻结资料、预算与模型版本。',
                            35:'集成与发布', 36:'三个分支已合入', 37:'不代替生产验收', 38:'集成结果见发布验收记录；模型与真人评价仍待实测。',
                            43:'复核与下一步', 44:'专业采用由使用者决定', 45:'事实、意见与材料缺口各自呈现。'}
                    for index, text in copy.items():
                        if index < len(texts):
                            texts[index].text = text
                if info.filename == 'ppt/slides/slide15.xml' and texts:
                    texts[-1].text = NOTICE
                    for paragraph in tree.iter():
                        if paragraph.tag.endswith('}p') and any(t.text == NOTICE for t in paragraph.iter() if t.tag.endswith('}t')):
                            for prop in paragraph.iter():
                                if prop.tag.endswith('}rPr'):
                                    prop.set('sz', '1400')
                body = ET.tostring(tree, encoding='utf-8', xml_declaration=True)
            elif info.filename == 'docProps/core.xml':
                tree = ET.fromstring(body)
                for element in tree.iter():
                    if element.tag.endswith('}creator') or element.tag.endswith('}lastModifiedBy'):
                        element.text = 'AuditTrace'
                body = ET.tostring(tree, encoding='utf-8', xml_declaration=True)
            output.writestr(info, body)
    # 包级重开仅证明结构可读，不替代PowerPoint渲染。
    with zipfile.ZipFile(target) as package:
        assert package.testzip() is None


def document(source, target, demo):
    doc = Document(source)
    paragraphs = list(doc.paragraphs)
    for table in doc.tables:
        paragraphs.extend(p for row in table.rows for cell in row.cells for p in cell.paragraphs)
    for section in doc.sections:
        paragraphs.extend(section.header.paragraphs)
        paragraphs.extend(section.footer.paragraphs)
    for paragraph in paragraphs:
        for run in paragraph.runs:
            run.text = run.text.replace('北京工商大学', '').replace('面向会计师事务所可持续信息鉴证项目组', '面向公开报告核对人员')
            run.text = run.text.replace('AI生成内容，用于披露关联核对和后续调查参考，不构成审计结论、鉴证结论或鉴证意见。', NOTICE).replace('系统生成内容用于披露关联核对和后续调查参考，不构成审计结论、鉴证结论或鉴证意见。', NOTICE)
    heading = doc.add_paragraph('本轮实现与验证状态', style='Heading 1')
    doc.add_paragraph(SCOPE)
    doc.add_paragraph('原方案的研究示例和拟建功能保留为设计依据；未绑定本轮运行的宝钢数字、版本链、知识增强及反证效果不能作为已验收成果。')
    doc.add_paragraph(f"当前本地任务{demo['task_id']}，运行{demo['run_id']}，为仅程序计算，外部模型调用0。正式复算{demo['overview']['program_levels']['formal']['value']}项，假设相容{demo['overview']['program_levels']['conditional']['value']}项，关系记录{demo['overview']['relation_total']}项；范围等前提不足仍保留。未完成人工专业复核，未部署到生产。")
    doc.add_paragraph('程序与模型对照使用17项冻结合成样例；13项现有可执行机器预期符合，4项待跨域能力验证。双人标注、真人节时和模型增益尚未测量，专业准确率保持空白。')
    doc.add_paragraph(NOTICE)
    for style in doc.styles:
        if style.type == 1:
            is_heading = style.name.startswith('Heading') or style.name == 'Title'
            style.font.name = '黑体' if is_heading else '宋体'
            style.font.size = Pt(16 if is_heading else 12)
            style.font.color.rgb = RGBColor(0, 0, 0)
            style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '黑体' if is_heading else '宋体')
    paragraphs.extend(doc.paragraphs[len(paragraphs):])
    for paragraph in list(doc.paragraphs) + [p for table in doc.tables for row in table.rows for cell in row.cells for p in cell.paragraphs]:
        heading = paragraph.style.name.startswith('Heading') or paragraph.style.name == 'Title'
        for run in paragraph.runs:
            run.font.name = '黑体' if heading else '宋体'
            run.font.size = Pt(16 if heading else 12)
            run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '黑体' if heading else '宋体')
    doc.core_properties.author = 'AuditTrace'
    doc.core_properties.last_modified_by = 'AuditTrace'
    doc.save(target)
    assert Document(target).paragraphs


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('source_root', type=Path)
    args = parser.parse_args()
    output = ROOT / 'outputs/ai4/materials'
    output.mkdir(parents=True, exist_ok=True)
    demo = json.loads((ROOT / 'outputs/ai4/demo/demo-index.json').read_text(encoding='utf-8'))
    presentation(args.source_root / '审迹智链(2).pptx', output / '答辩PPT_原模板修订候选_待渲染.pptx')
    document(args.source_root / '02_最终确定方案/审迹智链_年度报告与可持续发展报告关联核对_新方案书.docx', output / '方案书_原稿修订候选_待渲染.docx', demo)
    (output / '验收状态.json').write_text(json.dumps({'source_files_preserved': True, 'package_reopened': True,
            'office_render': 'blocked_wrong_desktop_session', 'submission_ready': False,
            'remaining': ['Office渲染逐页检查', '替换标注为历史业务的截图', '人工专业复核及发布门禁']}, ensure_ascii=False, indent=2), encoding='utf-8')
