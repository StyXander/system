"""从公开方案与完整40问生成统一宋体正文、黑体标题的Word文件。"""

from pathlib import Path
import hashlib
import json
import re

from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "决赛准备" / "项目同步发布_2026-10-14"
NOTICE = "AI生成内容，仅供审计计划阶段进一步核查，不构成审计结论或审计意见。"


def font(run, name, size):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), name)


def make(source, filename):
    text = source.read_text(encoding="utf-8")
    assert NOTICE in text
    assert not any(token in text for token in (r"\(", r"\)", r"\[", r"\]"))
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2.2)
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3", "Footer"):
        style = doc.styles[name]
        heading = name in ("Title", "Heading 1", "Heading 2", "Heading 3")
        style.font.name = "黑体" if heading else "宋体"
        style.font.size = Pt(16 if heading else 12)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), style.font.name)
        style.paragraph_format.line_spacing = 1.35
        style.paragraph_format.space_after = Pt(8 if heading else 6)
        if heading:
            style.paragraph_format.space_before = Pt(12)
            style.paragraph_format.keep_with_next = True
        for border in style._element.xpath(".//w:pBdr"):
            border.getparent().remove(border)
    # 中文标题和问答正文均由源稿驱动，避免另存一套不同的内容。
    for raw in text.splitlines():
        line = raw.strip()
        if not line:
            continue
        match = re.match(r"^(#{1,3})\s+(.+)", line)
        if match:
            depth, content = len(match[1]), match[2]
            style = "Title" if depth == 1 else f"Heading {depth - 1}"
            # 问题原题保留标点；其余标题采用纯文字。
            if not content.startswith("Q"):
                content = re.sub(r"[——：:·（）()]+", " ", content).strip()
            p = doc.add_paragraph(content, style)
            if depth == 1:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                font(run, "黑体", 16)
        else:
            content = re.sub(r"\[([^]]+)\]\(([^)]+)\)", r"\1（\2）", line)
            p = doc.add_paragraph(content)
            p.paragraph_format.widow_control = True
            p.paragraph_format.keep_with_next = line.startswith(("学生回答", "可能追问"))
            for run in p.runs:
                font(run, "宋体", 12)
                run.bold = line.startswith(("学生回答", "可能追问"))
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("第 "); font(run, "宋体", 12)
    field = OxmlElement("w:fldSimple"); field.set(qn("w:instr"), "PAGE"); footer._p.append(field)
    run = footer.add_run(" 页"); font(run, "宋体", 12)
    target = OUT / filename
    doc.save(target)
    reopened = Document(target)
    assert NOTICE in "\n".join(p.text for p in reopened.paragraphs)
    for p in reopened.paragraphs:
        heading = p.style.name in ("Title", "Heading 1", "Heading 2")
        for run in p.runs:
            assert run.font.size.pt == (16 if heading else 12)
            assert run._element.rPr.rFonts.get(qn("w:eastAsia")) == ("黑体" if heading else "宋体")
    return {"path": str(target.relative_to(ROOT)), "sha256": hashlib.sha256(target.read_bytes()).hexdigest(), "paragraph_count": len(reopened.paragraphs)}


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    (OUT / "核验").mkdir(exist_ok=True)
    answers = ROOT / "docs/审迹智链_答辩40问_2026-10-14.md"
    text = answers.read_text(encoding="utf-8")
    headings = re.findall(r"^### Q(\d+)\. (.+)$", text, re.M)
    assert [int(n) for n, _ in headings] == list(range(1, 41))
    assert text.count("学生回答") == 40
    files = [make(ROOT / "docs/审迹智链_决赛方案书_2026-10-14.md", "审迹智链_决赛项目方案书_2026-10-14.docx"),
             make(answers, "审迹智链_答辩40问_2026-10-14.docx")]
    (OUT / "核验/文档结构与字体核对.json").write_text(json.dumps({"question_count": 40, "question_order": "1-40", "body": "宋体12pt", "headings": "黑体16pt", "files": files}, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"question_count": 40, "files": files}, ensure_ascii=True))


if __name__ == "__main__":
    main()
