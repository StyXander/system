"""核验可重复的证据用途与资料撤回行为；这些断言不作审计有效性评分。"""

from copy import deepcopy
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from docx import Document

from backend.app.evidence_decisions import base_from_run, check_evidence, digest, investigate, load_scenarios
from backend.app.main import app


def question(investigation, qid):
    return next(q for q in investigation.questions if q["question_id"] == qid)


@pytest.mark.parametrize("case_id,step_index,qid,expected", [
    ("ANON-SALE-A", 0, "REV-CUT", "partial"),
    ("ANON-SALE-A", 1, "REV-CUT", "conflicting"),
    ("ANON-SALE-A", 2, "REV-CUT", "partial"),
    ("ANON-SALE-B", 0, "REV-ACC", "rule_supported"),
    ("ANON-SALE-B", 1, "REV-ACC", "conflicting"),
    ("ANON-AR-C", 0, "AR-VAL", "partial"),
    ("ANON-AR-C", 1, "AR-VAL", "rule_supported"),
    ("ANON-AR-C", 2, "AR-VAL", "partial"),
])
def test_guided_operations_use_registered_evidence_and_real_check(case_id, step_index, qid, expected):
    """首页操作仅登记资料选择；状态由真实接口计算，不能依赖前端示范文案。"""
    client = TestClient(app)
    catalog = client.get("/api/evidence-scenarios").json()
    scenario = next(s for s in catalog["scenarios"] if s["base_id"] == case_id)
    step = scenario["demonstration_steps"][step_index]
    original = scenario["investigation"]
    ids = {e["evidence_id"] for e in original["evidence_catalog"]}
    assert set(step["include_evidence_ids"] + step["exclude_evidence_ids"]) <= ids
    response = client.post("/api/evidence-checks", json={
        "base_kind": "validation_scenario", "base_id": case_id,
        "include_evidence_ids": step["include_evidence_ids"],
        "exclude_evidence_ids": step["exclude_evidence_ids"],
    })
    assert response.status_code == 200
    result = response.json()
    assert next(q for q in result["investigation"]["questions"] if q["question_id"] == qid)["state"] == expected
    assert result["provider_call_count"] == 0
    assert result["investigation"]["source_snapshot_sha256"] == original["source_snapshot_sha256"]
    assert result["investigation"]["professional_review_status"] == "not_performed"


@pytest.mark.parametrize("scenario_index", range(3))
@pytest.mark.parametrize("variant", range(8))
def test_anonymous_scope_and_withdrawal_variants(scenario_index, variant):
    base = deepcopy(load_scenarios()[scenario_index])
    original = digest(base)
    ids = [a["evidence_id"] for a in base["evidence"]]
    if variant == 0:
        checked = check_evidence(base, base_kind="validation_scenario", base_id=base["case_id"], include=[], exclude=[])
        assert not checked.changes
    elif variant == 1:
        full = investigate(base, ids)
        assert full.source_snapshot_sha256 == original
    elif variant == 2:
        empty = investigate(base, [])
        assert all(q["state"] in {"missing", "not_applicable"} for q in empty.questions)
    elif variant in {3, 4, 5, 6}:
        # 案例、客户、粒度、T0 任一边界失配都不能通过对应明细条件。
        key, value = {3: ("case_id", "OTHER"), 4: ("customer_id", "OTHER"), 5: ("granularity", "summary"), 6: ("available_at", "2027-01-01")}[variant]
        for atom in base["evidence"]:
            atom[key] = value
        inv = investigate(base, ids)
        assert all(q["state"] in {"missing", "not_applicable"} for q in inv.questions)
        assert not any(u["accepted"] for e in inv.evidence_decisions for u in e["uses"])
    else:
        first = deepcopy(base["evidence"][0]); first["evidence_id"] += "-DUP"
        base["evidence"].append(first)
        inv = investigate(base, [*ids, first["evidence_id"]])
        assert inv.evidence_decisions[-1]["duplicate"]
        assert all(first["evidence_id"] not in q["supporting_evidence_ids"] for q in inv.questions)
    if variant in {0, 1, 2}:
        assert digest(base) == original


@pytest.mark.parametrize("ticker", ["000858", "600938", "600309"])
@pytest.mark.parametrize("variant", range(4))
def test_real_report_summary_cannot_close_detail_conditions(ticker, variant):
    # 用真实企业标识核验汇总适配合同；金额与交易字段未编造。
    rows = [{"evidence_id": f"{ticker}-PUBLIC-FIELD", "field_name": "年报汇总收入", "source_id": f"{ticker}-REPORT"}]
    base = base_from_run({"case_id": ticker, "company_name": ticker, "t0": "2026-04-30"}, {"field_evidence": rows})
    selected = [] if variant == 1 else base["default_evidence_ids"]
    if variant == 2:
        base["evidence"][0]["kind"] = "receipt_record"
    if variant == 3:
        base["evidence"][0]["kind"] = "aging_schedule"
    inv = investigate(base, selected)
    assert len(inv.questions) == 5
    assert all(q["state"] == "missing" for q in inv.questions)
    assert inv.semantic_support_status == "not_verified"
    assert inv.professional_review_status == "not_performed"


def test_receipt_cannot_close_cutoff_and_acceptance_can_reveal_conflict():
    base = load_scenarios()[0]
    receipt = check_evidence(base, base_kind="validation_scenario", base_id=base["case_id"], include=["A-REC"], exclude=[])
    assert question(receipt.investigation, "REV-CUT")["state"] == "partial"
    conflict = check_evidence(base, base_kind="validation_scenario", base_id=base["case_id"], include=["A-ACC"], exclude=[])
    assert question(conflict.investigation, "REV-CUT")["state"] == "conflicting"
    assert question(conflict.investigation, "REV-OCC")["state"] == "rule_supported"
    withdrawn = check_evidence(base, base_kind="validation_scenario", base_id=base["case_id"], include=["A-ACC"], exclude=["A-CON"])
    assert question(withdrawn.investigation, "REV-CUT")["state"] == "partial"
    assert receipt.provider_call_count == 0


def test_amount_conflict_not_hidden_by_a_matching_invoice():
    base = load_scenarios()[1]
    supported = investigate(base, ["B-CON", "B-LED", "B-INV"])
    assert question(supported, "REV-ACC")["state"] == "rule_supported"
    conflicting = investigate(base, ["B-CON", "B-LED", "B-INV", "B-BAD"])
    assert question(conflicting, "REV-ACC")["state"] == "conflicting"


def test_summary_aging_and_partial_receipts_do_not_close_customer_conditions():
    base = load_scenarios()[2]
    inv = investigate(base, ["C-BAL", "C-REC", "C-SUM", "C-IMP"])
    assert question(inv, "AR-EXI")["state"] == "partial"
    assert "15000" in "".join(question(inv, "AR-EXI")["missing_conditions"])
    assert question(inv, "AR-VAL")["state"] == "partial"
    assert any(t["request_id"] == "AR-VAL:aging" for t in inv.material_requests)
    closed = investigate(base, ["C-BAL", "C-REC", "C-AGE", "C-IMP", "C-CON"])
    assert question(closed, "AR-EXI")["state"] == "rule_supported"
    assert closed.professional_review_status == "not_performed"


def test_selection_order_is_irrelevant_and_cross_base_references_rejected():
    base = load_scenarios()[1]
    a = check_evidence(base, base_kind="validation_scenario", base_id=base["case_id"], include=["B-BAD", "B-INV"], exclude=[])
    b = check_evidence(base, base_kind="validation_scenario", base_id=base["case_id"], include=["B-INV", "B-BAD", "B-INV"], exclude=[])
    assert a.model_dump() == b.model_dump()
    with pytest.raises(ValueError, match="白名单"):
        check_evidence(base, base_kind="validation_scenario", base_id=base["case_id"], include=["C-CON"], exclude=[])


def test_public_api_is_read_only_registered_only_and_document_uses_same_check(monkeypatch):
    import backend.app.main as main
    monkeypatch.setattr(main, "run_agent_chain", lambda **kwargs: pytest.fail("独立检验不能调用模型"))
    client = TestClient(app)
    catalog = client.get("/api/evidence-scenarios")
    assert catalog.status_code == 200
    assert len(catalog.json()["scenarios"]) == 3
    body = {"base_kind": "validation_scenario", "base_id": "ANON-SALE-A", "include_evidence_ids": ["A-ACC"]}
    response = client.post("/api/evidence-checks", json=body)
    assert response.status_code == 200
    payload = response.json()
    assert payload["provider_call_count"] == 0
    docx = client.post("/api/evidence-checks/report", json=body)
    assert docx.status_code == 200
    doc = Document(BytesIO(docx.content))
    texts = "\n".join(p.text for p in doc.paragraphs)
    assert payload["check_id"] in texts
    assert "未复核草稿" in texts
    for paragraph in doc.paragraphs:
        heading = paragraph.style.name.startswith("Heading") or paragraph.style.name == "Title"
        for run in paragraph.runs:
            assert run.font.size.pt == (16 if heading else 12)
            assert run.font.name == ("黑体" if heading else "宋体")
    assert client.post("/api/evidence-checks", json={**body, "free_text": "任意资料"}).status_code == 422
    assert client.post("/api/evidence-checks", json={**body, "exclude_evidence_ids": ["A-ACC"]}).status_code == 422
    assert client.post("/api/evidence-checks", json={**body, "base_id": "../../.env"}).status_code == 404


def test_private_run_authorization_is_required_before_check(monkeypatch):
    import backend.app.main as main
    from fastapi import HTTPException
    def denied(*args):
        raise HTTPException(403, "无权读取")
    monkeypatch.setattr(main, "get_run", denied)
    assert TestClient(app).post("/api/evidence-checks", json={"base_kind": "run", "base_id": "RUN-PRIVATE"}).status_code == 403


def test_registered_summary_preserves_customer_aging_request():
    from backend.app.main import _reconcile_registered_context
    from backend.app.schemas import RuleResult
    result = RuleResult(rule_id="R1", status="candidate", source_validation={}, metrics={}, risk_card={"data_gaps": ["客户级账龄明细未取得"], "requested_materials": ["客户级账龄明细", "期后回款"]})
    _reconcile_registered_context([result], [{"evidence_kind": "public_aging"}])
    assert result.risk_card["requested_materials"] == ["客户级账龄明细", "期后回款"]


def test_withdrawing_balance_removes_both_existence_and_valuation_precondition():
    base = load_scenarios()[2]
    all_ids = [a["evidence_id"] for a in base["evidence"]]
    complete = investigate(base, all_ids)
    assert question(complete, "AR-VAL")["state"] == "rule_supported"
    withdrawn = investigate(base, [eid for eid in all_ids if eid != "C-BAL"])
    assert question(withdrawn, "AR-EXI")["state"] == "partial"
    assert question(withdrawn, "AR-VAL")["state"] == "partial"
    assert any(task["request_id"] == "AR-VAL:balance" for task in withdrawn.material_requests)


def test_impossible_availability_and_mixed_valuation_basis_do_not_pass():
    from copy import deepcopy
    base = deepcopy(load_scenarios()[0])
    base["evidence"][2]["event_date"] = "2026-01-20"
    inv = investigate(base, [a["evidence_id"] for a in base["evidence"]])
    assert question(inv, "REV-CUT")["state"] == "partial"
    assert any("业务日期晚于" in u["reason"] for e in inv.evidence_decisions for u in e["uses"])
    base = deepcopy(load_scenarios()[2])
    base["evidence"][4]["currency"] = "USD"
    inv = investigate(base, [a["evidence_id"] for a in base["evidence"]])
    assert question(inv, "AR-VAL")["state"] == "partial"
    assert any("币种" in condition for condition in question(inv, "AR-VAL")["missing_conditions"])


def test_financial_industry_keeps_existing_specialized_scope():
    base = base_from_run({"case_id": "FINANCIAL-SCOPE", "industry_gate": {"industry_family": "banking"}}, {"field_evidence": []})
    inv = investigate(base)
    assert all(q["state"] == "not_applicable" for q in inv.questions)
    assert not inv.material_requests
    assert "行业专用规则" in inv.boundary


def test_transaction_receipt_can_support_customer_scope_with_balance_identity():
    base = deepcopy(load_scenarios()[2])
    base["evidence"][1]["granularity"] = "transaction"
    inv = investigate(base)
    existence = question(inv, "AR-EXI")
    assert "C-REC" in existence["supporting_evidence_ids"]
    assert existence["state"] == "partial"
    assert any("15000" in value for value in existence["missing_conditions"])
