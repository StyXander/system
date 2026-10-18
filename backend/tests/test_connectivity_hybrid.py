"""合成反例验证混合核验；离线提供方仅验证数据流，不证明模型专业能力。"""
from copy import deepcopy
from decimal import Decimal
from io import BytesIO
import json

from docx import Document
from fastapi import FastAPI
from fastapi.testclient import TestClient

from app.connectivity_conditions import upgrade_conditions
from app.connectivity_delivery import export_run, verify_export_same_source
from app.connectivity_engine import calculate
from app.connectivity_intelligence import enrich_calculation, integrate_analyses, refresh_analysis_status, targeted_context
from app.connectivity_pipeline import integrate_proposals, run_pipeline
from app.connectivity_schemas import (AdjustmentProposal, AIAnalysis, ConnectivityRun, DisclosureOccurrence, EvidenceRef,
                                    NarrativeClaim, ReportDocument, RunRequest, SearchCoverage)
from test_connectivity_engine import relationship
from test_connectivity_conditions_20261017 import row, run_with


def bridge_run():
    rows, check = relationship("EQ-001", ["1200", "1050"], ["left", "right"])
    for item in rows:
        item.label = "环保投入"
        item.year_status = "attributed"
        item.quality = "located"
        item.unit_source = "原表单位"
        item.context = item.context.replace("同一统计范围", "尚缺边界说明")
        item.context += f" 环保投入{item.value}万元。"
    rows.append(DisclosureOccurrence(occurrence_id="ADJ", document_id=rows[0].document_id, pdf_page=2,
        label="环保设备投入", value="150", unit="万元", currency="CNY", year="2025", year_status="attributed",
        quality="located", context="环保设备投入150万元，未说明是否已包含于环保治理投入。", topic="environment"))
    check.conditions[3].status = "unknown"
    check.adjustment_proposals = [AdjustmentProposal(occurrence_id="ADJ", label="环保设备投入", sign="+",
        reason="可能为两种统计范围之间的桥接项", evidence=[EvidenceRef(occurrence_id="ADJ", quote="环保设备投入150万元")],
        assumptions=["设备投入未重复计入治理投入"])]
    check = enrich_calculation(calculate(check, rows), rows)
    return ConnectivityRun(run_id="SYN-HYBRID", company="合成企业", period="2025", documents=[],
                           dataset_role="synthetic", occurrences=rows, checks=[check])


def opinion(run):
    c = run.checks[0]
    return AIAnalysis(check_id=c.check_id, verdict="plausible_explanation", summary="原文支持一个可能的范围调节路径，仍需核实是否重复计入。",
        reasons=["所列设备投入可在程序试算中桥接候选差额，但原文未确认两项互斥。"],
        supporting_evidence=[EvidenceRef(occurrence_id=run.occurrences[0].occurrence_id, quote="环保投入1200万元"),
                             EvidenceRef(occurrence_id="ADJ", quote="环保设备投入150万元")],
        assumptions=["设备投入未重复包含"], next_actions=["核实设备投资是否计入治理投入"])


def test_incomplete_conditions_still_have_bound_scenario_and_separate_opinion():
    run = bridge_run()
    check = run.checks[0]
    assert check.calculation["residual"] is None
    assert check.hypothetical["residual"] == "1500000"
    assert check.scenario_adjustment["remaining_residual"] == "0"
    assert check.calculation_status == "not_comparable"
    delta = integrate_analyses(run, [opinion(run)], {r.occurrence_id for r in run.occurrences}, "A3", "offline-test")
    assert delta["accepted_analysis_ids"] == [check.check_id]
    assert run.ai_analyses[0].analysis_quality == "medium"
    assert check.calculation_status == "not_comparable"


def test_adjustments_reject_missing_number_and_duplicate_cell():
    run = bridge_run()
    check = run.checks[0]
    check.adjustment_proposals[0].evidence[0].quote = "未说明是否已包含于环保治理投入"
    enrich_calculation(check, run.occurrences)
    assert not check.scenario_adjustment["applied"]
    assert "未包含调节额" in check.scenario_adjustment["rejected"][0]["reason"]
    run = bridge_run()
    check = run.checks[0]
    check.adjustment_proposals.append(check.adjustment_proposals[0].model_copy(deep=True))
    enrich_calculation(check, run.occurrences)
    assert len(check.scenario_adjustment["applied"]) == 1
    assert "重复" in check.scenario_adjustment["rejected"][0]["reason"]


def test_ai_opinion_accepts_computed_value_but_rejects_invented_value_and_quote():
    run = bridge_run()
    a = opinion(run)
    a.reasons.append("程序条件试算剩余差额为0，仍不代表范围已确认。")
    ids = {r.occurrence_id for r in run.occurrences}
    assert integrate_analyses(run, [a], ids, "A3", "offline")["accepted_analysis_ids"]
    b = opinion(run)
    b.summary = "原文显示还有99999万元，可能全部解释差异。"
    assert integrate_analyses(run, [b], ids, "A3", "offline")["analysis_rejections"]
    b = opinion(run)
    b.supporting_evidence[0].quote = "报告不存在的语句"
    assert integrate_analyses(run, [b], ids, "A3", "offline")["analysis_rejections"]


def test_opinion_becomes_stale_after_recomputation_and_exports_preserve_it():
    run = bridge_run()
    integrate_analyses(run, [opinion(run)], {r.occurrence_id for r in run.occurrences}, "A3", "offline")
    payload, _ = export_run(run, "json")
    restored = ConnectivityRun.model_validate(json.loads(payload))
    assert verify_export_same_source(restored, payload) == []
    assert restored.ai_analyses == run.ai_analyses
    restored.checks[0].conditions[3].reason = "新增人工范围说明"
    refresh_analysis_status(restored)
    assert restored.ai_analyses[0].status == "needs_refresh"
    csv, _ = export_run(run, "csv")
    assert "ai_analysis" in csv.decode() and "scenario_adjustment" in csv.decode()
    word, _ = export_run(run, "docx")
    text = "\n".join(p.text for p in Document(BytesIO(word)).paragraphs)
    assert run.ai_analyses[0].summary in text and "条件调节试算" in text
    tampered = json.loads(payload)
    tampered["ai_analyses"] = []
    assert verify_export_same_source(run, json.dumps(tampered).encode())


def test_update_keeps_prior_adjustments_conflicts_and_review():
    run = bridge_run()
    original = run.checks[0]
    original.review_status = "read_confirmed"
    original.conditions[0].status = "conflicting"
    original.conditions[0].reason = "主体原文冲突"
    fields = {"rule_id", "title", "topic", "members", "conditions", "relationship_evidence", "explanation", "explanation_evidence"}
    proposal = original.model_dump(include=fields)
    proposal["conditions"][0]["status"] = "supported"
    proposal["calculation_reason"] = "比较同一候选指标的两份披露"
    integrate_proposals(run, {"checks": [proposal]}, {r.occurrence_id for r in run.occurrences}, "A2")
    new = next(c for c in run.checks if c.check_id == original.check_id)
    assert new.review_status == "read_confirmed"
    assert new.conditions[0].status == "conflicting"
    assert new.adjustment_proposals and new.calculation_plan.status == "trial"


def test_different_metric_cannot_be_submitted_as_equal():
    run = bridge_run()
    fields = {"rule_id", "title", "topic", "members", "conditions", "relationship_evidence"}
    proposal = run.checks[0].model_dump(include=fields)
    proposal["mapping_type"] = "related_only"
    assert integrate_proposals(run, {"checks": [proposal]}, {r.occurrence_id for r in run.occurrences}, "A1")["rejected"]


def test_mismatching_total_is_discovered_and_complete_independently_of_arithmetic():
    rows = [row(1, "环保投入费用化", "60"), row(2, "环保投入资本化", "40"), row(3, "环保投入合计", "110")]
    from app.connectivity_pipeline import candidate_checks
    run = run_with(rows)
    checks = candidate_checks(rows, "2025", run=run)
    sums = [c for c in checks if c.operator == "sum"]
    assert sums and sums[0].calculation_status == "unexplained_difference"
    assert any(c.name == "组成完整" and c.status == "supported" for c in sums[0].conditions)


def test_total_growth_does_not_contradict_intensity_decrease():
    run = bridge_run()
    current = DisclosureOccurrence(occurrence_id="INT-NOW", document_id="ESG", pdf_page=3, label="单位产量排放强度", value="8",
        unit="吨二氧化碳当量/吨产品", year="2025", context="排放强度同比下降。单位产量排放强度8吨二氧化碳当量/吨产品。")
    previous = current.model_copy(update={"occurrence_id": "INT-PREV", "value": "10", "year": "2024", "context": "单位产量排放强度10吨二氧化碳当量/吨产品。"})
    run.occurrences += [current, previous]
    check = run.checks[0]
    check.narrative_claim = NarrativeClaim(indicator="单位产量排放强度", direction="decrease", current_id=current.occurrence_id,
        previous_id=previous.occurrence_id, evidence=[EvidenceRef(occurrence_id=current.occurrence_id, quote="排放强度同比下降")], basis_description="绑定单位产量强度，两期分母口径待核实")
    enrich_calculation(check, run.occurrences)
    assert check.narrative_test["direction_matches"]
    assert Decimal(check.narrative_test["change_percent"]) == Decimal("-20")


def test_targeted_material_keeps_member_ids_and_reports_budget_limit():
    run = bridge_run()
    rows, tasks = targeted_context(run, run.checks, limit=1)
    assert len(rows) == 1
    assert tasks[0]["essential_budget_exceeded"]
    assert "范围" in tasks[0]["missing_conditions"]
    assert all(ref in {r["occurrence_id"] for r in rows} for ref in tasks[0]["sent_ids"])


def test_three_role_pipeline_offline_end_to_end(monkeypatch, tmp_path):
    import app.connectivity_pipeline as pipeline
    run = bridge_run()
    docs = [ReportDocument(document_id=name, case_id="SYN-PAIR", company_id="SYN", company="合成企业", report_type=typ,
        period="2025", source_url="https://example.invalid/synthetic") for name, typ in [("SYNTHETIC-AR", "AR"), ("SYNTHETIC-ESG", "ESG")]]
    monkeypatch.setattr(pipeline, "load_registry", lambda root: docs)
    monkeypatch.setattr(pipeline, "extract_document", lambda root, doc, *a: ([r for r in run.occurrences if r.document_id == doc.document_id],
        SearchCoverage(document_id=doc.document_id, total_pages=2, parsed_pages=2, searched_pages=[1,2]), []))
    original = run.checks[0].model_copy(deep=True)
    original.adjustment_proposals = []
    monkeypatch.setattr(pipeline, "candidate_checks", lambda *a, **k: [deepcopy(original)])
    monkeypatch.setenv("DEEPSEEK_API_KEY", "offline-test-provider")
    monkeypatch.setenv("AUDITTRACE_DEMO_USE_EXTERNAL_MODEL", "true")
    seen = []

    def offline(role, payload):
        seen.append(role)
        if role == "A1":
            result = {"checks": []}
        elif role == "A2":
            assert payload["retrieval_tasks"]
            fields = {"rule_id", "title", "topic", "members", "conditions", "relationship_evidence"}
            p = original.model_dump(include=fields)
            p.update(mapping_type="same_metric", mapping_reason="两处环保投入候选，范围尚未核实", calculation_reason="按归一化数值比较候选差额",
                     adjustment_proposals=[run.checks[0].adjustment_proposals[0].model_dump()])
            result = {"checks": [p]}
        else:
            c = payload["candidate_checks"][0]
            calculated = next(item for item in payload["calculated_results"] if item["check_id"] == c["check_id"])
            assert calculated["scenario_adjustment"]["remaining_residual"] == "0"
            a = opinion(run).model_dump()
            a["check_id"] = c["check_id"]
            result = {"analyses": [a], "findings": []}
        return result, {"role": role, "model": "OFFLINE_TEST_DOUBLE", "status": "completed", "performed": True}

    monkeypatch.setattr(pipeline, "call_role", offline)
    result = run_pipeline(tmp_path, RunRequest(pair_id="SYN-PAIR"))
    assert seen == ["A1", "A2", "A3"]
    assert len(result.ai_analyses) == 1, result.agent_steps
    assert result.ai_analyses[0].status == "current"
    assert result.overview["ai_analysis"]["value"] == 1
    assert result.overview["scenario_bridges"]["value"] == 1
    assert result.checks[0].calculation["residual"] is None
    assert result.agent_steps[-1]["accepted_analysis_ids"]
    payload, _ = export_run(result, "json")
    restored = ConnectivityRun.model_validate_json(payload)
    assert restored.ai_analyses == result.ai_analyses
    assert verify_export_same_source(restored, payload) == []


def test_ai_review_saves_child_without_upgrading_computation(monkeypatch, tmp_path):
    from app import connectivity_api
    from test_connectivity_reviews import completed
    run = bridge_run()
    integrate_analyses(run, [opinion(run)], {r.occurrence_id for r in run.occurrences}, "A3", "offline")
    doc = ReportDocument(document_id="SYN", case_id="SYN-PAIR", company_id="SYN", company="合成企业", report_type="AR",
        period="2025", source_url="https://example.invalid/synthetic")
    monkeypatch.setattr(connectivity_api, "load_registry", lambda root: [doc])
    monkeypatch.setattr(connectivity_api, "demo_task_supabase_enabled", lambda: False)
    monkeypatch.setattr(connectivity_api, "run_pipeline", lambda *a, **k: run.model_copy(deep=True))
    app = FastAPI()
    app.include_router(connectivity_api.build_router(tmp_path, lambda r: None, lambda: None))
    with TestClient(app) as client:
        client.get("/api/connectivity/bootstrap")
        response = client.post("/api/connectivity/runs", json={"pair_id": "SYN-PAIR", "mode": "calculation_only"})
        task = completed(client, response.json()["task_id"])
        response = client.post(f'/api/connectivity/runs/{task["task_id"]}/reviews', json={"expected_revision": 1,
            "action": "ai_accept", "check_ids": [run.checks[0].check_id], "reason": "采纳为后续核查线索", "action_id": "synthetic-ai-review"})
        assert response.status_code == 200
        child = completed(client, response.json()["task_id"])["result"]
        assert child["ai_analyses"][0]["review_status"] == "accepted"
        assert child["checks"] == task["result"]["checks"]
        assert child["reviews"][-1]["before"]["review_status"] == "unreviewed"
        assert child["overview"]["ai_analysis"]["value"] == 1
        original = completed(client, task["task_id"])["result"]
        assert original["ai_analyses"][0]["review_status"] == "unreviewed"


def test_role_tool_contract_contains_opinions_and_bound_variables(monkeypatch):
    from app import connectivity_pipeline as pipeline
    monkeypatch.setenv("DEEPSEEK_API_KEY", "offline-tool-contract")
    monkeypatch.setenv("AUDITTRACE_DEMO_USE_EXTERNAL_MODEL", "true")
    contracts = []

    class Response:
        def __enter__(self):
            return self

        def __exit__(self, *args):
            return False

        def read(self):
            role = contracts[-1]["messages"][0]["content"]
            args = {"findings": [], "analyses": []} if "你负责独立综合研判" in role else {"checks": [], "notes": []}
            return json.dumps({"model": "OFFLINE_TOOL_CONTRACT", "choices": [{"message": {"tool_calls": [{"function": {
                "name": "submit_connectivity_relationships", "arguments": json.dumps(args)}}]}}], "usage": {"prompt_tokens": 1, "completion_tokens": 1}}).encode()

    def offline(request, **kwargs):
        contracts.append(json.loads(request.data))
        return Response()

    monkeypatch.setattr(pipeline, "urlopen", offline)
    run = bridge_run()
    payload = {"occurrences": [r.model_dump() for r in run.occurrences], "candidate_checks": [run.checks[0].model_dump()]}
    for role in ["A1", "A2", "A3", "single"]:
        pipeline.call_role(role, payload)
    for body in contracts:
        schema = body["tools"][0]["function"]["parameters"]
        assert schema["$defs"]["EvidenceRef"]["properties"]["occurrence_id"]["enum"] == [r.occurrence_id for r in run.occurrences]
    assert "AIAnalysis" in contracts[2]["tools"][0]["function"]["parameters"]["$defs"]
    assert "AdjustmentProposal" in contracts[1]["tools"][0]["function"]["parameters"]["$defs"]
