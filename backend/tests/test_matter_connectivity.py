"""事项主管线的合成对抗与接口检查；不计入真实企业或真实模型评价。"""
from datetime import date
import hashlib
import json
import threading
import time

import pymupdf
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient

from backend.app.connectivity_schemas import RunRequest, ConnectivityRun
from backend.app.connectivity_matter_schemas import (ReportSetRequest, AgentMatterResult, MatterAssessment,
    SustainabilityMatter, MatterRevisionRequest, InvestigationBudget)
from backend.app.connectivity_report_sets import make_report_set, create_report_set, resolve_documents
from backend.app.connectivity_matter_extraction import read_corpus, relevant_evidence, search_corpus
from backend.app.connectivity_matter_pipeline import run_matter_pipeline
from backend.app.connectivity_matters import integrate_result, apply_identification
from backend.app.connectivity_matter_revision import revise_matters


@pytest.mark.parametrize('text,expected', [
    ('在建工程达到预定可使用状态时，根据工程预算转入固定资产。', False),
    ('项目尚未正式投运。', False),
    ('预计2027年投产。', False),
    ('当工程完成转固时开始计提折旧。', False),
    ('节能改造项目已正式投运。', True),
    ('甲工厂项目于2025年完成转固。', True),
])
def test_policy_and_future_conditions_are_not_completed_project_events(text, expected):
    from backend.app.connectivity_matters import _reported_commissioning
    assert _reported_commissioning(text) is expected


@pytest.fixture
def reports(tmp_path, monkeypatch):
    monkeypatch.setenv('AUDITTRACE_CONNECTIVITY_CATALOG_PERSISTENCE','local')
    monkeypatch.setenv('AUDITTRACE_DEMO_TASK_PERSISTENCE','local')
    docs=[]
    directory=tmp_path/'backend/runtime/connectivity/sources';directory.mkdir(parents=True)
    for year in (2023,2024,2025):
        for kind in ('ESG','AR'):
            text=(f'{year}年合成企业 “绿色生产线改造项目”位于甲工厂。\n气候变化和极端天气可能影响产量。\n计划到2027年实现能耗下降目标。'
                if kind=='ESG' else f'{year}年合成企业财务报表附注\n固定资产使用寿命与减值测试\n现金流预测考虑产品需求变化。\n并未单独说明气候因素的估值影响。')
            pdf=pymupdf.open();page=pdf.new_page();page.insert_text((50,70),text,fontname='china-s',fontsize=11)
            raw=pdf.tobytes();pdf.close();sha=hashlib.sha256(raw).hexdigest();(directory/(sha+'.pdf')).write_bytes(raw)
            docs.append({'document_id':f'{kind}-{year}','case_id':f'SYN-{year}','company_id':'SYN','company':'合成企业','report_type':kind,
                'period':str(year),'disclosure_date':f'{year+1}-04-30','source_url':'https://example.invalid/source.pdf','file_sha256':sha,'pages':1})
    (tmp_path/'backend/connectivity_reports.json').write_text(json.dumps({'documents':docs},ensure_ascii=False))
    report_set=create_report_set(tmp_path,ReportSetRequest(document_ids=[d['document_id'] for d in docs],as_of_date=date(2026,10,26)))
    return tmp_path,report_set


def request(report_set,**extra):
    return RunRequest(report_set_id=report_set.report_set_id,mode='five_agents',budget=InvestigationBudget(max_matters=1),**extra)


def fake_provider(role,payload):
    data={'matter_id':payload['matter_id'],'input_revision':payload['input_revision'],'summary':role+'合成测试回执'}
    esg=[e for e in payload['evidence'] if e['report_type']=='ESG']
    ar=[e for e in payload['evidence'] if e['report_type']=='AR']
    if role=='A1':
        data['matters']=[{'name':'绿色生产线改造项目','matter_type':'transition_project','subject':'合成企业','location':'甲工厂',
            'identity_basis':'原文均明确相同项目名与甲工厂','evidence_ids':[e['evidence_id'] for e in esg]}]
    elif role=='A2' and not payload.get('tool_results'):
        data['tool_requests']=[{'tool_name':'search_matter_reports','arguments':{'query':'使用寿命 减值 现金流预测','report_type':'AR'},'reason':'补查估计假设'}]
    elif role=='A2':
        data['paths']=[{'path_id':'candidate','matter_id':payload['matter_id'],'financial_item':'固定资产减值',
            'steps':[{'label':'绿色改造事项','explanation':'原文提到具体项目','origin':'report_statement','support':'supported','evidence_ids':[esg[0]['evidence_id']]},
                     {'label':'影响现金流假设的候选机制','explanation':'尚无项目到估计假设的直接解释','origin':'model_hypothesis','support':'supported','evidence_ids':[ar[0]['evidence_id']]}]}]
    elif role=='A3':
        data['events']=[{'event_id':'candidate','matter_id':payload['matter_id'],'kind':'promise','label':'绿色项目目标',
            'report_period':esg[0]['report_period'],'occurred_period':'','evidence_ids':[esg[0]['evidence_id']]}]
    elif role=='A4':
        data['explanations']=[{'hypothesis_id':'candidate','matter_id':payload['matter_id'],'explanation':'可能未单独披露而已汇总',
            'discriminating_evidence':'需要甲工厂项目台账与资产组资料','supporting_evidence_ids':[],'status':'supported_as_reported'}]
    elif role=='A5':
        data['assessment']={'matter_id':payload['matter_id'],'summary':'项目存在，财务路径仅部分支持，缺少具体归属。',
            'relation_support':'supported','amount_comparability':'comparable','supporting_evidence_ids':[esg[0]['evidence_id']],
            'unresolved':['项目到资产组假设尚缺直接证据']}
    return AgentMatterResult.model_validate(data),{'performed':True,'request_attempted':True,'status':'completed','model':'synthetic-test-provider','input_tokens':10,'output_tokens':5}


def test_multi_year_set_and_legacy_request_contract(reports):
    root,rs=reports
    assert rs.periods==['2023','2024','2025']
    assert request(rs).analysis_scope=='matter_connectivity'
    assert RunRequest(pair_id='SYN-2025').mode=='three_agents'
    assert len(resolve_documents(root,request(rs))[1])==6
    with pytest.raises(ValueError):RunRequest(pair_id='SYN-2025',report_set_id=rs.report_set_id)


def test_set_accepts_one_report_and_reports_gaps(reports):
    root,rs=reports
    one=create_report_set(root,ReportSetRequest(document_ids=['AR-2023'],as_of_date=date(2026,10,26)))
    assert len(one.document_ids)==1 and one.coverage_gaps==['2023年缺少ESG报告']


def test_as_of_rejects_future_or_missing_disclosure(reports):
    root,rs=reports
    with pytest.raises(ValueError):create_report_set(root,ReportSetRequest(document_ids=['AR-2025'],as_of_date=date(2025,1,1)))


def test_no_amount_climate_narrative_is_indexed(reports):
    root,rs=reports;_,docs=resolve_documents(root,request(rs))
    corpus,coverage=read_corpus(root,docs)
    assert any('极端天气' in e.quote for e in corpus)
    seed,meta=relevant_evidence(corpus)
    assert len({b.report_period for b in seed})==3 and not meta['seed_truncated']
    assert all(c['total_pages']==1 for c in coverage)


def test_bm25_alias_filters_and_honest_semantic_status(reports):
    root,rs=reports;_,docs=resolve_documents(root,request(rs));corpus,_=read_corpus(root,docs)
    result=search_corpus(corpus,'减值',report_type='AR',periods=['2024'])
    assert result['evidence'] and {e['report_period'] for e in result['evidence']}=={'2024'}
    assert result['semantic_status']=='not_configured'
    assert not search_corpus(corpus,'完全不存在的XYZABC')['evidence']


def test_five_roles_tool_feedback_and_supported_cap(reports):
    root,rs=reports;checkpoints=[]
    run=run_matter_pipeline(root,request(rs),caller=fake_provider,checkpoint=lambda r:checkpoints.append(len(r.provider_calls)))
    assert set(t['role'] for t in run.matter_tasks)=={'A1','A2','A3','A4','A5'}
    assert len(run.provider_calls)==6 and run.matter_tool_receipts
    assert run.matter_paths[0].support=='limited'
    assert run.matter_assessments[0].relation_support=='limited'
    assert run.matter_assessments[0].amount_comparability=='needs_evidence'
    assert run.matter_explanations[0].status=='hypothesis'
    assert max(checkpoints)==6
    assert run.matter_overview['pending']>0


def test_parallel_workers_share_one_call_budget(reports):
    root,rs=reports;calls=[];guard=threading.Lock()
    def caller(role,payload):
        with guard:calls.append(role)
        time.sleep(.01)
        return fake_provider(role,payload)
    req=request(rs);req.budget.max_provider_calls=3
    run=run_matter_pipeline(root,req,caller=caller)
    assert len(calls)==len(run.provider_calls)==3
    assert run.versions['budget']['remaining']==0
    assert any(t['status']=='failed' for t in run.matter_tasks)
    assert run.model_status=='five_agents_partial'


@pytest.mark.parametrize('persistent_error', [False, True])
def test_schema_repair_is_bounded_and_preserves_failure_receipt(reports, persistent_error):
    from backend.app.connectivity_matter_agents import MatterCallFailure
    root,rs=reports;attempts=[]
    def caller(role,payload):
        assert payload['investigation_round']['remaining_tool_rounds'] >= 0
        if role=='A3':
            attempts.append(payload)
            if persistent_error or len(attempts)==1:
                raise MatterCallFailure({'failure_code':'MODEL_SCHEMA_INVALID','status':'failed',
                    'performed':True,'request_attempted':True,'retryable':False,
                    'invalid_fields':[{'location':['tool_arguments'],'type':'json_invalid'}]})
        return fake_provider(role,payload)
    run=run_matter_pipeline(root,request(rs),caller=caller)
    assert len(attempts)==2
    assert attempts[1]['output_validation_feedback']['code']=='MODEL_SCHEMA_INVALID'
    receipts=[c for c in run.provider_calls if c['role']=='A3']
    assert len(receipts)==2 and receipts[0]['failure_code']=='MODEL_SCHEMA_INVALID'
    assert receipts[-1]['status']==('failed' if persistent_error else 'completed')
    assert next(t for t in run.matter_tasks if t['role']=='A3')['status']==('failed' if persistent_error else 'completed')


@pytest.mark.parametrize('ignores_finalization', [False,True])
def test_last_tool_feedback_gets_one_bounded_final_synthesis(reports, ignores_finalization):
    root,rs=reports;calls=[]
    def caller(role,payload):
        output,record=fake_provider(role,payload)
        if role=='A1':
            calls.append(payload.copy())
            if not payload.get('finalize_only') or ignores_finalization:
                from backend.app.connectivity_matter_schemas import MatterToolCall
                output.tool_requests=[MatterToolCall(tool_name='lookup_ontology_rule',arguments={'keyword':'减值'},reason='核查专业规则')]
        return output,record
    run=run_matter_pipeline(root,request(rs),caller=caller)
    assert len(calls)==3 and calls[-1]['finalize_only']
    task=next(t for t in run.matter_tasks if t['role']=='A1')
    assert task['status']==('partial' if ignores_finalization else 'completed')
    assert len([c for c in run.provider_calls if c['role']=='A1'])==3


def test_unsupported_commitment_year_is_rejected_then_corrected(reports):
    from backend.app.connectivity_matter_schemas import MatterCommitment
    root,rs=reports;attempts=[]
    def caller(role,payload):
        output,record=fake_provider(role,payload)
        if role=='A3':
            attempts.append(payload.copy())
            if len(attempts)==1:
                output.commitments=[MatterCommitment(commitment_id='invalid-year',matter_id=payload['matter_id'],
                    label='没有原文支持的年份',proposed_period='2025',target_year='2099',evidence_ids=[payload['evidence'][0]['evidence_id']])]
        return output,record
    run=run_matter_pipeline(root,request(rs),caller=caller)
    assert len(attempts)==2 and attempts[-1]['evidence_validation_feedback']
    assert not any(c.target_year=='2099' for c in run.matter_commitments)
    assert any('承诺年份没有原文支持' in r.get('reason','') for r in run.matter_rejections)
    assert next(t for t in run.matter_tasks if t['role']=='A3')['status']=='completed'


def test_missing_key_never_creates_fake_role_completion(reports,monkeypatch):
    root,rs=reports;monkeypatch.delenv('DEEPSEEK_API_KEY',raising=False)
    run=run_matter_pipeline(root,request(rs))
    assert run.model_status=='five_agents_partial' and not run.provider_calls
    assert all(t['status']=='skipped' for t in run.matter_tasks)


def test_planning_groups_uninvestigated_candidates_instead_of_inventing_findings(reports):
    from backend.app.connectivity_assurance_planning import build_assurance_planning
    root,rs=reports;req=request(rs);req.mode='calculation_only'
    run=run_matter_pipeline(root,req)
    planning=build_assurance_planning(run)
    assert len(run.matters)>1
    assert len(planning['risk_findings'])==1
    assert planning['risk_findings'][0]['matter_ids']==[]
    assert '初筛候选尚未深查' in planning['risk_findings'][0]['reason']
    investigated=run_matter_pipeline(root,request(rs),caller=fake_provider)
    planning=build_assurance_planning(investigated)
    assert {mid for row in planning['risk_findings'] for mid in row['matter_ids']}=={a.matter_id for a in investigated.matter_assessments}


@pytest.mark.parametrize('bad_kind',['foreign_ref','wrong_revision','wrong_matter','invented_period'])
def test_role_proposals_reject_identity_and_evidence_errors(reports,bad_kind):
    root,rs=reports;run=run_matter_pipeline(root,RunRequest(report_set_id=rs.report_set_id,mode='calculation_only'))
    matter=run.matters[0];eid=matter.evidence_ids[0]
    value={'matter_id':matter.matter_id,'input_revision':run.result_revision,'events':[{'event_id':'x','matter_id':matter.matter_id,
        'kind':'promise','label':'目标','report_period':'2023','evidence_ids':[eid]}]}
    if bad_kind=='foreign_ref':value['events'][0]['evidence_ids']=['NO-SUCH-EVIDENCE']
    if bad_kind=='wrong_revision':value['input_revision']=999
    if bad_kind=='wrong_matter':value['matter_id']='OTHER'
    if bad_kind=='invented_period':value['events'][0]['occurred_period']='1991'
    with pytest.raises(ValueError):integrate_result(run,matter,'A3',AgentMatterResult.model_validate(value),run.result_revision)


def test_human_split_preserves_lineage_and_rejects_all_evidence_split(reports):
    root,rs=reports;run=run_matter_pipeline(root,request(rs),caller=fake_provider)
    m=next(m for m in run.matters if m.identity_status=='evidence_linked')
    body=MatterRevisionRequest(expected_revision=1,action_id='split-operation',action='split',matter_ids=[m.matter_id],evidence_ids=m.evidence_ids,reason='测试拆分理由')
    with pytest.raises(ValueError):revise_matters(run,body)
    body.evidence_ids=m.evidence_ids[:1];revise_matters(run,body)
    children=[m2 for m2 in run.matters if m.matter_id in m2.lineage]
    assert len(children)==2 and all(c.investigation_status=='pending' for c in children)
    assert not any(p.matter_id==m.matter_id for p in run.matter_paths)


def test_export_uses_saved_matter_records(reports):
    from backend.app.connectivity_delivery import export_run
    from docx import Document
    from io import BytesIO
    root,rs=reports;run=run_matter_pipeline(root,request(rs),caller=fake_provider)
    payload,_=export_run(run,'json');assert len(json.loads(payload)['matters'])==len(run.matters)
    csv_data,_=export_run(run,'csv');assert b'matter_paths' in csv_data
    word,_=export_run(run,'docx');text=''.join(p.text for p in Document(BytesIO(word)).paragraphs)
    assert run.run_id in text and run.matters[0].name in text


def test_matter_api_owner_export_and_refresh(reports):
    from backend.app.connectivity_api import build_router
    root,rs=reports;app=FastAPI();app.include_router(build_router(root,lambda _:None,lambda:None))
    with TestClient(app) as client:
        client.get('/api/connectivity/bootstrap')
        created=client.post('/api/connectivity/runs',json={'report_set_id':rs.report_set_id,'mode':'calculation_only'},headers={'Idempotency-Key':'matter-api-idempotent'})
        assert created.status_code==200,created.text
        tid=created.json()['task_id']
        for _ in range(100):
            current=client.get('/api/connectivity/runs/'+tid)
            if current.json()['status'] not in {'queued','running'}:break
            time.sleep(.02)
        assert current.json()['result']['matter_schema_version']
        assert client.get(f'/api/connectivity/runs/{tid}/matters').status_code==200
        assert client.get(f'/api/connectivity/runs/{tid}/export/json').status_code==200
        with TestClient(app) as other:
            other.get('/api/connectivity/bootstrap')
            assert other.get('/api/connectivity/runs/'+tid).status_code==403
            assert other.get(f'/api/connectivity/runs/{tid}/export/json').status_code==403


def test_single_user_public_report_selection_and_five_agent_run_without_login(reports,monkeypatch):
    from backend.app import connectivity_matter_pipeline as pipeline, main
    from backend.app.connectivity_api import build_router
    from backend.app.connectivity_catalog_api import build_catalog_router
    root,rs=reports
    monkeypatch.setenv('AUDITTRACE_SINGLE_USER','true')
    monkeypatch.setenv('AUDITTRACE_PUBLIC_DEMO','true')
    monkeypatch.setenv('AUDITTRACE_DEMO_EXECUTOR_MODE','web')
    monkeypatch.setenv('AUDITTRACE_MODEL_DAILY_INPUT_TOKENS','1')
    monkeypatch.setenv('AUDITTRACE_MODEL_DAILY_RUN_LIMIT','1')
    def forbidden():raise AssertionError('single-user requests must not consult the old quota ledger')
    monkeypatch.setattr(main,'_public_model_ledger',forbidden)
    original=pipeline.run_matter_pipeline
    monkeypatch.setattr(pipeline,'run_matter_pipeline',lambda *a,**k:original(*a,**k,caller=fake_provider))
    app=FastAPI();app.include_router(build_router(root,lambda r:main._enforce_public_model_quota(r,force=True),forbidden));app.include_router(build_catalog_router(root))
    with TestClient(app) as client:
        client.get('/api/connectivity/bootstrap')
        me=client.get('/api/auth/me').json()
        assert not me['authentication_required'] and me['can_import'] and not me['authenticated']
        selected=client.post('/api/connectivity/report-sets',json={'document_ids':rs.document_ids,'as_of_date':'2026-10-26'})
        assert selected.status_code==200,selected.text
        # 连续两次运行，低至1的旧日次数/token变量也不阻断演示。
        for _ in range(2):
            created=client.post('/api/connectivity/runs',json={'report_set_id':selected.json()['report_set_id'],'mode':'five_agents','budget':{'max_matters':1}})
            assert created.status_code==200,created.text
            tid=created.json()['task_id']
            for _ in range(200):
                current=client.get('/api/connectivity/runs/'+tid).json()
                if current['status'] not in {'queued','running'}:break
                time.sleep(.02)
            assert current['result']['dataset_role']=='synthetic'
            assert {t['role'] for t in current['result']['matter_tasks']}=={'A1','A2','A3','A4','A5'}
            assert current['result']['provider_calls']
        # 同一个人的另一浏览器和另一个部署可直接读、导出并继续整理历史事项。
        with TestClient(app) as second:
            opened=second.get('/api/connectivity/runs/'+tid)
            assert opened.status_code==200 and opened.json()['can_modify']
            assert second.get(f'/api/connectivity/runs/{tid}/export/json').status_code==200
            assert tid in [r['task_id'] for r in second.get('/api/connectivity/runs').json()['items']]


def test_identity_merge_rejects_generic_note_mixed_with_project(reports):
    from backend.app.connectivity_matter_schemas import MatterProposal
    root,rs=reports;run=run_matter_pipeline(root,RunRequest(report_set_id=rs.report_set_id,mode='calculation_only'))
    esg=next(e for e in run.matter_evidence if e.report_type=='ESG');ar=next(e for e in run.matter_evidence if e.report_type=='AR')
    before=[m.matter_id for m in run.matters]
    apply_identification(run,[MatterProposal(name='绿色生产线改造项目',matter_type='transition_project',identity_basis='主题相近',evidence_ids=[esg.evidence_id,ar.evidence_id])])
    assert [m.matter_id for m in run.matters]==before and run.matter_rejections


def test_role_rejection_is_atomic_and_cannot_forge_human_review(reports):
    root,rs=reports;run=run_matter_pipeline(root,request(rs),caller=fake_provider);matter=run.matters[0]
    count=len(run.matter_paths)
    result=AgentMatterResult(matter_id=matter.matter_id,paths=[{'path_id':'new','matter_id':matter.matter_id,'financial_item':'新科目',
        'steps':[{'label':'新机制','origin':'human_review','evidence_ids':matter.evidence_ids,'support':'supported'}]}],
        relations=[{'edge_id':'bad','matter_id':matter.matter_id,'source_id':'missing','target_id':matter.matter_id,'relation_kind':'component_of'}])
    with pytest.raises(ValueError):integrate_result(run,matter,'A2',result,1)
    assert len(run.matter_paths)==count
    result=AgentMatterResult(matter_id=matter.matter_id,assessment={'matter_id':matter.matter_id,'summary':'模型意见','review_status':'accepted'})
    integrate_result(run,matter,'A5',result,1)
    assert next(a for a in run.matter_assessments if a.matter_id==matter.matter_id).review_status=='unreviewed'


def test_page_tool_returns_citable_matter_evidence_and_rejects_arbitrary_arithmetic(reports):
    from backend.app.connectivity_tool_registry import execute_tool,ToolRequest
    root,rs=reports;run=run_matter_pipeline(root,RunRequest(report_set_id=rs.report_set_id,mode='calculation_only'))
    page=execute_tool(root,run,ToolRequest('open_source_page',{'document_id':'ESG-2023','page':1}))
    assert page['status']=='hit' and page['result']['evidence'][0]['source_verified']
    attempt=execute_tool(root,run,ToolRequest('recompute_decimal',{'left':'100','right':'200','operator':'sum'}))
    assert attempt['status']=='rejected' and attempt['result'] is None


def test_numeric_tool_uses_registered_original_and_exposes_empty_coverage(reports):
    from backend.app.connectivity_tool_registry import execute_tool,ToolRequest
    root,rs=reports;run=run_matter_pipeline(root,RunRequest(report_set_id=rs.report_set_id,mode='calculation_only'))
    receipt=execute_tool(root,run,ToolRequest('locate_matter_numbers',{'document_id':'AR-2023','pages':[1]}))
    assert receipt['status'] in {'hit','no_match'} and receipt['result']['coverage']['total_pages']==1
    rejected=execute_tool(root,run,ToolRequest('locate_matter_numbers',{'document_id':'unknown'}))
    assert rejected['status']=='rejected'


def test_many_to_many_relation_endpoints_survive_graph_projection(reports):
    from backend.app.connectivity_matters import finish_projection
    root,rs=reports;run=run_matter_pipeline(root,request(rs),caller=fake_provider);m,n=run.matters[:2]
    result=AgentMatterResult(matter_id=m.matter_id,relations=[{'edge_id':'shared','matter_id':m.matter_id,'source_id':m.matter_id,'target_id':n.matter_id,
        'relation_kind':'component_of','assertion_origin':'model_hypothesis','evidence_status':'supported','evidence_ids':m.evidence_ids}])
    integrate_result(run,m,'A2',result,1);finish_projection(run)
    edge=next(e for e in run.matter_graph['edges'] if e.get('id')=='shared')
    assert edge['source']==m.matter_id and edge['target']==n.matter_id and edge['status']=='limited'


def test_source_changes_invalidate_matter_assessment_without_overwriting_history(reports):
    from backend.app.connectivity_graph import refresh_run_sources
    from backend.app.connectivity_extraction import source_path
    root,rs=reports;run=run_matter_pipeline(root,request(rs),caller=fake_provider)
    doc=next(d for d in run.documents if d.document_id==run.matter_evidence[0].document_id)
    source_path(root,doc).write_bytes(b'changed')
    changed=refresh_run_sources(root,run)
    assert any(e.invalid_reason for e in changed.matter_evidence)
    assert changed.versions['source_validity']['status']=='invalid'
    assert not any(e.invalid_reason for e in run.matter_evidence)


def test_target_revision_preserves_old_commitment_and_original_time(reports):
    from backend.app.connectivity_matter_schemas import MatterCommitment
    root,rs=reports;run=run_matter_pipeline(root,RunRequest(report_set_id=rs.report_set_id,mode='calculation_only'));m=run.matters[0]
    eid=m.evidence_ids[0]
    run.matter_commitments=[MatterCommitment(commitment_id='old',matter_id=m.matter_id,label='原目标',proposed_period='2023',target_year='2027',evidence_ids=[eid])]
    result=AgentMatterResult(matter_id=m.matter_id,commitments=[{'commitment_id':'new','matter_id':m.matter_id,'label':'新目标候选','evidence_ids':[eid],
        'proposed_period':'2023','target_year':'2027','supersedes':'old'}])
    integrate_result(run,m,'A3',result,1)
    assert next(c for c in run.matter_commitments if c.commitment_id=='old').status=='revised'
    assert next(c for c in run.matter_commitments if c.commitment_id=='new').status=='not_due'
    assert all(not e.occurred_period for e in run.matter_events)


def test_scanned_pdf_is_reported_as_coverage_gap_not_no_risk(tmp_path):
    from backend.app.connectivity_schemas import ReportDocument
    directory=tmp_path/'backend/runtime/connectivity/sources';directory.mkdir(parents=True)
    pdf=pymupdf.open();pdf.new_page();raw=pdf.tobytes();pdf.close();sha=hashlib.sha256(raw).hexdigest();(directory/(sha+'.pdf')).write_bytes(raw)
    doc=ReportDocument(document_id='scan',case_id='scan',company_id='SYN',company='合成扫描',report_type='ESG',period='2025',source_url='https://example.invalid',file_sha256=sha)
    corpus,coverage=read_corpus(tmp_path,[doc])
    assert not corpus and coverage[0]['failed_pages'] and coverage[0]['scanned_pages']==[1]


def test_restart_keeps_matter_checkpoint_and_real_role_order(reports):
    from backend.app.demo_run_tasks import DemoRunTaskStore
    root,rs=reports;run=run_matter_pipeline(root,request(rs),caller=fake_provider)
    directory=root/'task-restart';first=DemoRunTaskStore(directory,task_id_prefix='CONNJOB')
    task=first._new_task(rs.report_set_id,{'kind':'connectivity','analysis_scope':'matter_connectivity','mode':'five_agents','owner':'a'})
    assert set(task['agent_steps'])=={'A1','A2','A3','A4','A5'}
    task.update(status='running',result=run.model_dump(mode='json'));first._save_file(task);first.shutdown()
    second=DemoRunTaskStore(directory,task_id_prefix='CONNJOB')
    try:
        snapshot=second.snapshot(task['task_id'])
        assert snapshot['status']=='interrupted'
        assert snapshot['result']['run_id']==run.run_id and snapshot['result']['provider_calls']
    finally:second.shutdown()


@pytest.mark.parametrize('ablation,omitted',[('no_temporal','A3'),('no_counterevidence','A4')])
def test_ablation_does_not_fake_omitted_role(reports,ablation,omitted):
    root,rs=reports;req=request(rs);req.ablation=ablation
    run=run_matter_pipeline(root,req,caller=fake_provider)
    assert not any(c['role']==omitted for c in run.provider_calls)
    assert any(t['role']==omitted and t['status']=='skipped' for t in run.matter_tasks)
    assert run.versions['ablation']==ablation and run.model_status=='five_agents_partial'


def test_matter_rollback_switch_refuses_new_work_without_relabeling(reports,monkeypatch):
    from backend.app.connectivity_api import build_router
    root,rs=reports;monkeypatch.setenv('CONNECTIVITY_MATTER_ENABLED','false')
    app=FastAPI();app.include_router(build_router(root,lambda _:None,lambda:None))
    with TestClient(app) as client:
        boot=client.get('/api/connectivity/bootstrap').json()
        assert boot['default_mode']=='three_agents' and boot['matter_enabled'] is False
        response=client.post('/api/connectivity/runs',json={'report_set_id':rs.report_set_id,'mode':'five_agents'})
        assert response.status_code==503
