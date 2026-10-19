"""AI4交付合同回归，合成输入不冒充企业或真人评价。"""
import hashlib
import json
from datetime import datetime, timezone, timedelta
from io import BytesIO

import pymupdf
import pytest
from docx import Document
from fastapi.testclient import TestClient

from app import connectivity_bootstrap as bootstrap
from app.connectivity_delivery import export_run
from app.connectivity_facts import execution_provenance, project_facts
from app.connectivity_overview import build_overview
from app.connectivity_schemas import ConnectivityRun, ReportDocument, SearchCoverage, ExtractionExclusion
from app.demo_run_tasks import DemoRunTaskStore, SupabaseDemoRunTaskStore, result_expiry_iso
from test_connectivity_reviews import workspace, completed


def test_history_session_isolation_pagination_and_no_payload(workspace):
    client, app, original = workspace
    first = client.get('/api/connectivity/runs?limit=1').json()
    assert first['items'][0]['task_id'] == original['task_id']
    assert 'result' not in first['items'][0]
    assert 'owner' not in json.dumps(first)
    with TestClient(app) as stranger:
        assert stranger.get('/api/connectivity/runs').status_code == 409
        stranger.get('/api/connectivity/bootstrap')
        assert stranger.get('/api/connectivity/runs').json()['items'] == []
        # 明确保持“已知链接公开，列表私有”的既有分享合同。
        assert stranger.get('/api/connectivity/runs/' + original['task_id']).status_code == 200
    assert client.get('/api/connectivity/runs?limit=21').status_code == 422
    assert client.get('/api/connectivity/runs?offset=-1').status_code == 422


def test_new_connectivity_result_retention_and_provenance(workspace):
    client, _, original = workspace
    expires = datetime.fromisoformat(original['result_expires_at'])
    assert expires > datetime.now(timezone.utc) + timedelta(days=29)
    exported = client.get(f"/api/connectivity/runs/{original['task_id']}/export/json").json()
    assert exported['overview']['execution_provenance']['kind'] == 'synthetic'
    assert exported['overview']['execution_provenance'] == original['result']['overview']['execution_provenance']


def test_minimum_retention_does_not_relabel_old_results(monkeypatch):
    monkeypatch.setenv('AUDITTRACE_DEMO_TASK_RETENTION_HOURS', '1')
    expiry = datetime.fromisoformat(result_expiry_iso(minimum_seconds=30*86400))
    assert expiry > datetime.now(timezone.utc) + timedelta(days=29)
    assert datetime.fromisoformat(result_expiry_iso()) < datetime.now(timezone.utc) + timedelta(hours=2)


def test_local_history_restart_expired_and_other_owner(tmp_path):
    store = DemoRunTaskStore(tmp_path, task_id_prefix='CONNJOB')
    try:
        task = store._new_task('PAIR', {'kind': 'connectivity', 'owner': 'a'})
        task.update(status='completed', result={'private': True}, result_expires_at='2020-01-01T00:00:00+00:00')
        store._save_file(task)
        assert store.list_owned('b') == []
        assert store.list_owned('a')[0]['status'] == 'expired'
        assert store.list_owned('a')[0]['result'] is None
    finally:
        store.shutdown()


def test_remote_history_ownership_filter_and_projection():
    class Client:
        def _request(self, method, path, **kwargs):
            assert kwargs['params']['request_payload->>owner'] == 'eq.' + 'a'*64
            assert 'result' not in kwargs['params']['select'].split(',')
            assert kwargs['params']['offset'] == '5'
            return [{'task_id': 'one', 'request_payload': {'owner': 'a'*64}, 'status': 'completed', 'result_expires_at': '2020-01-01T00:00:00+00:00'},
                    {'task_id': 'two', 'request_payload': {'owner': 'b'}}]
    store = SupabaseDemoRunTaskStore(Client())
    try:
        rows = store.list_owned('a'*64, offset=5)
        assert [r['task_id'] for r in rows] == ['one']
        assert rows[0]['status'] == 'expired'
    finally:
        store.shutdown()


@pytest.mark.parametrize('state', ['missing', 'hash_mismatch', 'page_count_mismatch', 'ready'])
def test_source_pair_readiness(tmp_path, monkeypatch, state):
    pdf = pymupdf.open()
    pdf.new_page().insert_text((50, 50), 'source')
    content = pdf.tobytes()
    pdf.close()
    digest = hashlib.sha256(content).hexdigest()
    docs = [ReportDocument(document_id=kind, case_id='PAIR', company_id='S', company='合成', report_type=kind,
             period='2025', source_url='https://example.invalid', file_sha256=digest,
             pages=2 if state == 'page_count_mismatch' else 1) for kind in ['AR', 'ESG']]
    monkeypatch.setattr(bootstrap, 'load_registry', lambda _: docs)
    path = bootstrap.source_path(tmp_path, docs[0])
    path.parent.mkdir(parents=True)
    if state != 'missing':
        path.write_bytes(b'broken' if state == 'hash_mismatch' else content)
    result = bootstrap.source_readiness(tmp_path)
    assert result['ready'] == (state == 'ready')
    assert result['pairs'][0]['documents'][0]['status'] == state
    assert result['remote_status'] == 'not_probed'
    if state == 'ready':
        path.write_bytes(b'changed after readiness')
        changed = bootstrap.source_readiness(tmp_path)
        assert not changed['ready']
        assert changed['pairs'][0]['documents'][0]['status'] == 'hash_mismatch'


def test_export_summary_bands_counts_and_word_fonts(workspace):
    _, _, original = workspace
    raw = original['result'].copy()
    raw.pop('execution_provenance', None)
    raw.pop('ai_generated_content_notice', None)
    raw.pop('personal_observation_count', None)
    run = ConnectivityRun.model_validate(raw)
    overview = build_overview(run)
    assert sum(c['value'] for c in overview['summary_bands'].values()) == len(run.checks)
    assert overview['coverage_counts']['failed_pages']['unit'] == '页'
    data, _ = export_run(run, 'docx')
    document = Document(BytesIO(data))
    assert '可确认' in '\n'.join(p.text for p in document.paragraphs)
    for paragraph in document.paragraphs:
        heading = (paragraph._p.style or '').startswith('Heading') or paragraph._p.style == 'Title'
        for text_run in paragraph.runs:
            assert text_run.font.size.pt == (16 if heading else 12)
            assert text_run.font.name == ('黑体' if heading else '宋体')


def test_no_calls_is_not_real_model_success():
    run = ConnectivityRun(run_id='S', company='合成', period='2025', documents=[], execution_mode='three_agents', model_status='three_agents_completed')
    assert execution_provenance(run)['kind'] == 'unverified_model'


def test_same_ip_four_attempts_keep_budget_and_concurrency(tmp_path):
    from app.public_model import PublicModelLedger, QuotaConfig, PublicModelQuotaError
    ledger = PublicModelLedger(tmp_path, config=QuotaConfig(per_ip=2, daily_runs=3, max_concurrent=1))
    first = ledger.reserve('same-ip')
    with pytest.raises(PublicModelQuotaError, match='并发'):
        ledger.reserve('other-ip')
    ledger.settle(first, input_tokens=1, output_tokens=1)
    outcomes = ['accepted']
    for _ in range(3):
        try:
            reservation = ledger.reserve('same-ip')
            ledger.settle(reservation, input_tokens=1, output_tokens=1)
            outcomes.append('accepted')
        except PublicModelQuotaError:
            outcomes.append('limited')
    assert outcomes == ['accepted', 'accepted', 'limited', 'limited']
    third = ledger.reserve('other-ip')
    ledger.settle(third, input_tokens=1, output_tokens=1)
    with pytest.raises(PublicModelQuotaError, match='今日'):
        ledger.reserve('third-ip')
