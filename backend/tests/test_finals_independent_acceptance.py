"""独立验收反例。仅调用本地函数和测试桩，不调用外部模型，不修改被验收源码。"""
from copy import deepcopy
from pathlib import Path
import sys

SOURCE = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(SOURCE))

from backend.app.connectivity_claim_taxonomy import classify_claim
from backend.app.connectivity_financial_bridge import build
from backend.app.connectivity_graph import build_impact_graph, invalidate_graph
from backend.app.connectivity_temporal import compare_periods
from backend.app.connectivity_peer_benchmark import peer_benchmark
from backend.app import connectivity_agent_orchestrator as agent
from backend.tests.test_financial_bridge import _docs, _pages
from backend.tests.test_agent_orchestrator import _plan, _revision
from backend.app.connectivity_schemas import ConnectivityRun


def test_future_reduction_target_must_not_be_realized():
    result = classify_claim('公司目标到2030年减少碳排放20%。')
    assert result['planned_or_realized'] != 'realized', result


def test_completed_output_is_a_reachable_ninth_category():
    result = classify_claim('报告期内建成污水处理设施，新增处理产能10万吨。')
    assert result['claim_kind'] == 'output', result


def test_changed_annual_report_invalidates_its_dependent_edges():
    graph = build_impact_graph(build(None, documents=_docs(), pages=_pages(), industry='steel'))
    financial_ids = {n['id'] for n in graph['nodes'] if n['document_id'] == 'TEST_AR'}
    dependencies = [e['id'] for e in graph['edges'] if e['target'] in financial_ids]
    result = invalidate_graph(deepcopy(graph), {'TEST_AR': 'f' * 64})
    marked = {e['id'] for e in result['edges'] if e.get('invalid')}
    assert dependencies and set(dependencies) <= marked, {'dependencies': dependencies, 'invalidated': sorted(marked)}


def test_unrelated_reverse_sentence_must_not_weaken_both_alternatives():
    revision = _revision()
    receipt = {'status': 'hit', 'output_digest': 'Q-1', 'result': {
        'quote': '公司员工培训覆盖率未达到目标；本页同时披露绿色项目正常建设。'}}
    transitions = agent._gate_transitions(revision, {'edge_id': 'EDGE-1'}, receipt, [], 'INV-R01')
    assert not any(t['after'] == 'weakened' for t in transitions), transitions


def test_challenger_competing_explanation_enters_live_ledger(monkeypatch):
    monkeypatch.setattr(agent, 'execute_tool', lambda *a: {
        'tool_name': 'open_source_page', 'args_hash': 'a', 'status': 'hit',
        'output_digest': 'Q-1', 'source_hash': 's', 'elapsed_ms': 0,
        'result': {'quote': '绿色项目预算披露。'}, 'reason': ''})
    revision = _revision()
    run = ConnectivityRun(run_id='RUN-AUDIT', company='测试钢铁', period='2025', documents=[])
    def challenger(payload):
        return {'hypotheses': [{'edge_id': 'EDGE-1', 'explanation': '该金额可能为分阶段预算而非已发生支出。'}]}, {}
    result = agent.run_investigation(None, run, bridge={'edges': []}, revision=revision,
                                     plan=_plan(count=1), challenger=challenger)
    candidates = [t['hypothesis_id'] for t in result['transitions'] if t['before'] == '']
    live = {h['hypothesis_id'] for h in revision['hypotheses']}
    assert candidates and set(candidates) <= live, {'new_transition_ids': candidates, 'live_ledger_ids': sorted(live)}


def test_final_plan_is_updated_for_ui_consumer(monkeypatch):
    monkeypatch.setattr(agent, 'execute_tool', lambda *a: {
        'tool_name': 'open_source_page', 'args_hash': 'a', 'status': 'hit',
        'output_digest': 'Q-1', 'source_hash': 's', 'elapsed_ms': 0,
        'result': {'quote': '预算原文。'}, 'reason': ''})
    plan = _plan(count=2)
    run = ConnectivityRun(run_id='RUN-AUDIT', company='测试钢铁', period='2025', documents=[])
    result = agent.run_investigation(None, run, bridge={'edges': []}, revision=_revision(),
                                     plan=plan, challenger=None, max_rounds=1)
    assert plan['best_path'][0]['action_id'] != 'ACT-1' or result.get('final_plan'), {
        'caller_plan_top1': plan['best_path'][0]['action_id'],
        'last_event': result['events'][-1], 'result_keys': list(result)}


def test_same_period_is_not_cross_year_growth():
    current = {'label': '环保投入', 'value': '12', 'unit': '亿元',
               'document_id': 'DOC-CURRENT', 'pdf_page': 1, 'period': '2025'}
    prior = {**current, 'value': '10', 'document_id': 'DOC-OTHER'}
    result = compare_periods(current, prior)
    assert result['available'] is False, result


def test_restatement_or_boundary_change_blocks_direct_yoy():
    current = {'label': '环保投入', 'value': '12', 'unit': '亿元', 'company': '甲公司',
               'document_id': 'DOC-2025', 'pdf_page': 1, 'period': '2025', 'scope': '合并',
               'restated': True, 'definition_changed': True}
    prior = {**current, 'value': '10', 'document_id': 'DOC-2024', 'period': '2024', 'scope': '母公司',
             'restated': False, 'definition_changed': False}
    result = compare_periods(current, prior)
    assert result['available'] is False, result


def test_incompatible_peer_years_and_scopes_must_not_rank():
    registry = {'authorized': True, 'samples': [
        {'company': '甲', 'label': '环保投入', 'value': '10', 'unit': '亿元',
         'period': '2023', 'scope': '母公司', 'industry': '食品', 'source': '甲2023 p1'},
        {'company': '乙', 'label': '环保投入', 'value': '20', 'unit': '亿元',
         'period': '2024', 'scope': '合并', 'industry': '钢铁', 'source': '乙2024 p2'},
        {'company': '丙', 'label': '环保投入', 'value': '30', 'unit': '亿元',
         'period': '2025', 'scope': '国内子公司', 'industry': '科技', 'source': '丙2025 p3'}]}
    result = peer_benchmark('环保投入', '15', '亿元', registry=registry)
    assert result['available'] is False, result


def test_ui_consumes_new_graph_and_agent_contracts():
    script = (SOURCE / 'assets/official-v4/connectivity-graph.js').read_text(encoding='utf-8')
    assert 'key_path' in script and 'investigation' in script, '前端未消费key_path和investigation'


def test_scenario_has_two_editable_controls():
    script = (SOURCE / 'assets/official-v4/connectivity-graph.js').read_text(encoding='utf-8')
    assert 'type="range"' in script or 'type="number"' in script, '情景仅有静态预计算展示，无参数输入'


def test_docx_consumes_scenario_v2_values():
    script = (SOURCE / 'backend/app/connectivity_delivery.py').read_text(encoding='utf-8')
    assert 'scenario.get("baseline"' in script and 'scenario.get("scenarios"' in script, 'DOCX仍只遍历旧ranges合同'


def test_docx_actual_export_keeps_scenario_results():
    from io import BytesIO
    from docx import Document
    from backend.app.connectivity_delivery import export_run
    from backend.app.connectivity_financial_scenario import build_scenario
    bridge = build(None, documents=_docs(), pages=_pages(), industry='steel')
    scenario = build_scenario(bridge)
    scenario['baseline']['value'] = '123456.78'
    run = ConnectivityRun(run_id='RUN-AUDIT', company='测试钢铁', period='2025', documents=_docs(),
                          financial_claims=bridge['claims'], financial_nodes=bridge['nodes'],
                          financial_edges=bridge['edges'], impact_graph=build_impact_graph(bridge),
                          financial_scenario=scenario)
    payload, _ = export_run(run, 'docx')
    document = Document(BytesIO(payload))
    text = '\n'.join([p.text for p in document.paragraphs] +
                     [c.text for t in document.tables for r in t.rows for c in r.cells])
    assert '123456.78' in text, '实际Word导出缺少scenario.baseline.value，复现不落盘Word'


def test_carbon_emissions_cannot_be_used_as_energy_cost_base():
    from backend.app.connectivity_financial_scenario import build_scenario
    from backend.app.connectivity_financial_bridge import _amounts
    bridge = build(None, documents=_docs(), pages=_pages(), industry='steel')
    text = '减少碳排放64.5万吨二氧化碳当量'
    amounts = _amounts(text)
    assert amounts and amounts[0]['dimension'] == 'co2', amounts
    bridge['claims'] = [{'claim_id': 'C-CO2', 'document_id': 'TEST_ESG', 'pdf_page': 1,
                         'text': text, 'claim_kind': 'performance', 'amounts': amounts}]
    result = build_scenario(bridge)
    assert result['available'] is False, result
