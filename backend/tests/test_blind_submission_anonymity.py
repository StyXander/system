"""盲审 DOCX 匿名性的回归测试：删掉显示节点不等于删掉了图片。

对应 2026-09-08 独立验收 A01：上一版脚本只摘除 `w:drawing` 节点，图片和它的关系
仍留在压缩包里，评委解压即可还原导师签名。这里用合成文档同时锁住两件事：
只删节点必须被判为失败；连关系一起删才算通过。
"""

from __future__ import annotations

import importlib.util
import zipfile
import zlib
from io import BytesIO
from pathlib import Path
from typing import Any

from docx import Document

ROOT = Path(__file__).resolve().parents[2]
SPEC = importlib.util.spec_from_file_location("blind_copy_module", ROOT / "scripts" / "build_blind_submission_copy.py")
blind = importlib.util.module_from_spec(SPEC)  # type: ignore[arg-type]
assert SPEC and SPEC.loader
SPEC.loader.exec_module(blind)  # type: ignore[attr-defined]


def _png(seed: int) -> bytes:
    """现场构造一张合法的 1×1 PNG；像素不同，才不会被 python-docx 按内容去重成一部件。"""

    def chunk(kind: bytes, data: bytes) -> bytes:
        return len(data).to_bytes(4, "big") + kind + data + zlib.crc32(kind + data).to_bytes(4, "big")

    raw = bytes((0, seed % 251, (seed * 7) % 251, (seed * 13) % 251))
    return (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", (1).to_bytes(4, "big") + (1).to_bytes(4, "big") + bytes((8, 2, 0, 0, 0)))
        + chunk(b"IDAT", zlib.compress(raw))
        + chunk(b"IEND", b"")
    )


def _doc_with_pictures(tmp_path: Path, count: int, name: str) -> Path:
    """生成一份把图片放在表格单元格里的合成文档，模拟实名母版。"""

    document = Document()
    table = document.add_table(rows=1, cols=max(count, 1))
    for index in range(count):
        paragraph = table.cell(0, index).add_paragraph()
        paragraph.add_run().add_picture(BytesIO(_png(index + 1)))
    path = tmp_path / name
    document.save(path)
    return path


def _media_members(path: Path) -> list[str]:
    with zipfile.ZipFile(path) as archive:
        return sorted(n for n in archive.namelist() if n.startswith("word/media/"))


def test_only_stripping_the_drawing_leaves_the_image_in_the_package(tmp_path: Path) -> None:
    """复现 A01：删节点但不删关系时，媒体部件仍在，必须被复验判为失败。"""

    source = _doc_with_pictures(tmp_path, 1, "named.docx")
    document = Document(source)
    removed, rids = blind._strip_drawings(document.tables[0].cell(0, 0))  # type: ignore[attr-defined]
    assert removed == 1 and rids
    stripped_only = tmp_path / "stripped-only.docx"
    document.save(stripped_only)

    assert _media_members(stripped_only) == _media_members(source), "旧行为：图片仍打包在盲审件内"
    report = blind._verify_anonymity(source, stripped_only, [])  # type: ignore[attr-defined]
    assert report["passed"] is False
    assert report["orphan_media_parts"] and report["dangling_image_relations"]


def test_dropping_unreferenced_relations_removes_the_media(tmp_path: Path) -> None:
    """修好之后：孤立关系被删，媒体部件不再写入，复验通过。"""

    source = _doc_with_pictures(tmp_path, 1, "named.docx")
    document = Document(source)
    _, rids = blind._strip_drawings(document.tables[0].cell(0, 0))  # type: ignore[attr-defined]
    dropped = blind._drop_unreferenced_parts(document, rids)  # type: ignore[attr-defined]
    assert dropped == sorted(rids)
    fixed = tmp_path / "fixed.docx"
    document.save(fixed)

    assert _media_members(fixed) == []
    report = blind._verify_anonymity(source, fixed, dropped)  # type: ignore[attr-defined]
    assert report["passed"] is True, report
    assert report["orphan_media_parts"] == []
    assert report["dangling_image_relations"] == []


def test_referenced_images_are_kept(tmp_path: Path) -> None:
    """正文里仍在使用的图不能被顺手删掉，否则匿名会破坏版面。"""

    source = _doc_with_pictures(tmp_path, 2, "named-two.docx")
    assert len(_media_members(source)) == 2
    document = Document(source)
    _, rids = blind._strip_drawings(document.tables[0].cell(0, 0))  # type: ignore[attr-defined]
    dropped = blind._drop_unreferenced_parts(document, rids)  # type: ignore[attr-defined]
    kept = tmp_path / "kept.docx"
    document.save(kept)

    remaining = _media_members(kept)
    assert len(remaining) == 1, "另一张正文图必须保留"
    assert blind._verify_anonymity(source, kept, dropped)["passed"] is True  # type: ignore[attr-defined]


def test_identity_text_in_headers_and_metadata_is_detected(tmp_path: Path) -> None:
    """身份线索扫描要覆盖页眉页脚与元数据，而不只看正文文本层。"""

    document = Document()
    document.add_paragraph("正常正文内容。")
    section = document.sections[0]
    header_paragraph = section.header.paragraphs[0]
    header_paragraph.text = "指导教师：刘金钊 北京工商大学"
    document.core_properties.author = "刘金钊"
    source = tmp_path / "named.docx"
    document.save(source)

    report: dict[str, Any] = blind._verify_anonymity(source, source, [])  # type: ignore[attr-defined]
    assert report["passed"] is False
    assert "word/header1.xml" in report["identity_text_hits_by_part"]
    assert report["core_properties"]["author"] == "刘金钊"


def test_normalize_target_resolves_relative_and_parent_paths() -> None:
    """关系目标解析覆盖同目录与上级目录两种写法。"""

    assert blind._normalize_target("word/document.xml", "media/image1.jpeg") == "word/media/image1.jpeg"  # type: ignore[attr-defined]
    assert blind._normalize_target("word/header1.xml", "../customXml/item1.xml") == "customXml/item1.xml"  # type: ignore[attr-defined]
