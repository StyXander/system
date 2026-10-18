"""混合推理业务反例；离线模型替身只验证流程，不充当真实模型效果。"""
from copy import deepcopy
from decimal import Decimal
from io import BytesIO
import json

from docx import Document

from app.ai_opinion import build_program_hints, program_hint
from app.connectivity_delivery import export_run, verify_export_same_source
from app.connectivity_engine import calculate
from app.connectivity_extraction import PARSER_CACHE_VERSION, retrieve_targeted_pages
from app.connectivity_intelligence import integrate_analyses, refresh_analysis_status, pack_role_context
from app.connectivity_schemas import AIAnalysis, DisclosureOccurrence, EvidenceRef, ExplanationHypothesis, ReportDocument, RunRequest, SearchCoverage
from app.judgment_layer import Materiality
from test_connectivity_hybrid import bridge_run, opinion


def test_ratio_thresholds_keep_distinct_precision():
    thresholds = Materiality.from_base(Decimal('1'))
    assert thresholds.overall > thresholds.performance > thresholds.clearly_trivial > 0
    assert thresholds.performance == Decimal('0.0075')
    assert thresholds.clearly_trivial == Decimal('0.0005')


def test_unknown_evidence_does_not_become_high_enterprise_risk_or_force_review():
    run = bridge_run()
    check = run.checks[0]
    check.calculation = {'left': '100', 'right': '100', 'normalized_unit': '元'}
    hint = program_hint(check)
    assert hint['priority'] == 'low'
    assert hint['coverage']['unknown']
    assert 'needs_human' not in hint and 'confidence' not in hint and 'score' not in hint


def test_model_changes_own_verdict_and_priority_without_overwriting_program_facts():
    run = bridge_run()
    before = deepcopy(run.checks[0].calculation)
    item = opinion(run)
    item.priority = 'low'
    item.priority_reason = '设备投资提供具体解释路径，优先查明是否重复包含'
    integrate_analyses(run, [item], {o.occurrence_id for o in run.occurrences}, 'A3', 'offline')
    hint = build_program_hints(run)[0]
    assert hint['program_priority'] == 'medium'
    assert hint['priority'] == 'low' and hint['priority_source'] == 'model'
    assert run.ai_analyses[0].verdict == 'plausible_explanation'
    assert run.checks[0].calculation == before


def test_equivalent_unit_display_and_partial_invalid_advice_keep_valid_opinion():
    run = bridge_run()
    item = opinion(run)
    item.reasons.append('程序候选差额为150万元，可在假设下用设备投入桥接。')
    item.next_actions.append('请确认不存在的99999万元是否已纳入。')
    delta = integrate_analyses(run, [item], {o.occurrence_id for o in run.occurrences}, 'A3', 'offline')
    assert delta['accepted_analysis_ids']
    assert delta['analysis_rejections']
    assert any('150万元' in reason for reason in run.ai_analyses[0].reasons)
    assert not any('99999' in text for text in run.ai_analyses[0].next_actions)


def test_text_without_numbers_still_gets_independent_analysis():
    run = bridge_run()
    check = run.checks[0]
    check.operator = 'narrative'
    check.rule_id = 'NAR-001'
    check.calculation = {'statement': '加强环保投入'}
    check.hypothetical = None
    check.calculation_status = 'not_applicable'
    item = opinion(run)
    item.summary = '加强环保投入主要描述行动力度，现有文字没有明确承诺金额同比上升。'
    item.reasons = ['应先识别该主张是否涉及金额变化，不能直接把行动措辞转成增长承诺。']
    delta = integrate_analyses(run, [item], {o.occurrence_id for o in run.occurrences}, 'A3', 'offline')
    assert delta['accepted_analysis_ids'] and run.ai_analyses
    assert program_hint(check)['difference']['absolute'] is None


def test_counterevidence_retracts_hypothesis_and_preserves_prior_opinion():
    run = bridge_run()
    ids = {o.occurrence_id for o in run.occurrences}
    first = opinion(run)
    first.hypotheses = [ExplanationHypothesis(hypothesis_id='scope', explanation='设备投入可能构成未纳入的范围差异', status='plausible',
        supporting_evidence=first.supporting_evidence, assumptions=['设备未重复包含'])]
    integrate_analyses(run, [first], ids, 'A3', 'offline')
    counter = DisclosureOccurrence(occurrence_id='COUNTER', document_id=run.occurrences[0].document_id, pdf_page=3,
        label='组成说明', context='设备投资已包含于所列环保治理投入。', quality='text_only')
    run.occurrences.append(counter)
    second = opinion(run)
    second.verdict = 'potential_conflict'
    second.summary = '新增说明表明设备投资已经包含，原调节假设不能继续作为解释。'
    second.reasons = ['已包含的设备投资不能再次从差额扣除，范围解释应撤回。']
    second.counter_evidence = [EvidenceRef(occurrence_id='COUNTER', quote=counter.context)]
    second.hypotheses = [first.hypotheses[0].model_copy(update={'status': 'refuted', 'counter_evidence': second.counter_evidence})]
    second.revision_note = '新增组成说明否定原来的未纳入假设'
    integrate_analyses(run, [second], ids | {'COUNTER'}, 'A3', 'offline')
    assert run.checks[0].hypotheses[0].status == 'refuted'
    assert run.ai_analyses[0].verdict == 'potential_conflict'
    assert run.analysis_history[0]['previous']['verdict'] == 'plausible_explanation'


def test_targeted_retrieval_reaches_cached_pages_outside_initial_candidates(tmp_path):
    run = bridge_run()
    doc_id = run.occurrences[0].document_id
    document = ReportDocument(document_id=doc_id, case_id='SYN', company_id='SYN', company='合成企业', report_type='AR',
        period='2025', file_sha256='a'*64, source_url='https://example.invalid/synthetic')
    run.documents = [document]
    run.checks[0].hypotheses = [ExplanationHypothesis(hypothesis_id='coverage', explanation='设备统计范围可能解释剩余差额', retrieval_queries=['运营控制权'])]
    directory = tmp_path/'backend/runtime/connectivity/extraction'
    directory.mkdir(parents=True)
    page = {'page': 70, 'width': 595, 'height': 842, 'error': '', 'lines': [
        {'text': '运营控制权口径的统计范围包括受控制的生产设施。', 'side': 0, 'bbox': [10,20,500,40]}]}
    (directory/f'{document.file_sha256}-{PARSER_CACHE_VERSION}.json').write_text(json.dumps([page]), encoding='utf-8')
    tasks = retrieve_targeted_pages(tmp_path, run, run.checks)
    added = [o for o in run.occurrences if '-SUP-' in o.occurrence_id]
    assert added and added[0].pdf_page == 70 and added[0].value is None
    assert tasks[0]['scanned_pages'][doc_id] == [70]
    assert tasks[0]['new_evidence_ids'] == [added[0].occurrence_id]


def offline_pipeline(monkeypatch, run):
    from app import connectivity_pipeline as pipeline
    documents = [ReportDocument(document_id=name, case_id='SYN-PAIR', company_id='SYN', company='合成企业', report_type=kind,
        period='2025', source_url='https://example.invalid/synthetic') for name, kind in [('SYNTHETIC-AR','AR'),('SYNTHETIC-ESG','ESG')]]
    run.documents = documents
    monkeypatch.setattr(pipeline, 'load_registry', lambda root: documents)
    monkeypatch.setattr(pipeline, 'extract_document', lambda root, doc, *args: ([o for o in run.occurrences if o.document_id==doc.document_id],
        SearchCoverage(document_id=doc.document_id,total_pages=3,parsed_pages=3,searched_pages=[1,2,3]), []))
    monkeypatch.setattr(pipeline, 'candidate_checks', lambda *a,**k: deepcopy(run.checks))
    monkeypatch.setenv('DEEPSEEK_API_KEY','OFFLINE_TEST_ONLY')
    monkeypatch.setenv('AUDITTRACE_DEMO_USE_EXTERNAL_MODEL','true')
    return pipeline


def test_failed_a2_batch_does_not_block_a3_or_erase_results(monkeypatch, tmp_path):
    run=bridge_run()
    pipeline=offline_pipeline(monkeypatch,run)
    seen=[]
    def call(role,payload):
        seen.append(role)
        if role=='A2':
            raise ValueError('MODEL_NETWORK_UNAVAILABLE')
        result={'checks':[]} if role=='A1' else {'findings':[], 'analyses':[opinion(run).model_dump()]}
        return result, {'role':role,'performed':True,'status':'completed','model':'OFFLINE_TEST_DOUBLE'}
    monkeypatch.setattr(pipeline,'call_role',call)
    result=pipeline.run_pipeline(tmp_path,RunRequest(pair_id='SYN-PAIR'))
    assert seen == ['A1','A2','A3']
    assert result.model_status=='three_agents_partial'
    assert result.ai_analyses and result.checks[0].hypothetical
    assert any(c['status']=='failed' for c in result.provider_calls)
    assert not result.ai_opinions and result.program_hints


def test_single_model_new_relation_is_analysed_in_same_run(monkeypatch,tmp_path):
    run=bridge_run()
    pipeline=offline_pipeline(monkeypatch,run)
    monkeypatch.setattr(pipeline,'candidate_checks',lambda *a,**k: [])
    seen=[]
    def call(role,payload):
        assert role=='single'
        seen.append(payload.get('analysis_only',False))
        if not payload.get('analysis_only'):
            fields=set(pipeline.RelationshipProposal.model_fields)
            result={'checks':[run.checks[0].model_dump(include=fields)],'analyses':[]}
        else:
            item=opinion(run).model_dump()
            item['check_id']=payload['candidate_checks'][0]['check_id']
            result={'checks':[],'analyses':[item]}
        return result, {'role':role,'performed':True,'status':'completed','model':'OFFLINE_TEST_DOUBLE'}
    monkeypatch.setattr(pipeline,'call_role',call)
    result=pipeline.run_pipeline(tmp_path,RunRequest(pair_id='SYN-PAIR',mode='single_model'))
    assert seen==[False,True]
    assert result.ai_analyses[0].check_id==result.checks[0].check_id
    assert result.model_status=='single_completed'


def test_export_typography_and_program_facts_remain_same_source():
    run=bridge_run()
    integrate_analyses(run,[opinion(run)],{o.occurrence_id for o in run.occurrences},'A3','offline')
    run.program_hints=build_program_hints(run)
    payload,_=export_run(run,'json')
    assert verify_export_same_source(run,payload)==[]
    word,_=export_run(run,'docx')
    document=Document(BytesIO(word))
    paragraphs=[*document.paragraphs,*[p for t in document.tables for r in t.rows for c in r.cells for p in c.paragraphs]]
    for paragraph in paragraphs:
        heading=paragraph.style.name.startswith('Heading') or paragraph.style.name=='Title'
        for item in paragraph.runs:
            if item.text:
                assert item.font.name==('黑体' if heading else '宋体')
                assert item.font.size.pt==(16 if heading else 12)
    tampered=json.loads(payload)
    tampered['program_hints']=[]
    assert verify_export_same_source(run,json.dumps(tampered).encode())


def test_incremental_retry_only_analyzes_missing_relationship(monkeypatch,tmp_path):
    run=bridge_run()
    pipeline=offline_pipeline(monkeypatch,run)
    def complete(role,payload):
        result={'checks':[]} if role!='A3' else {'findings':[], 'analyses':[opinion(run).model_dump()]}
        return result, {'role':role,'performed':True,'status':'completed','model':'OFFLINE_TEST_DOUBLE'}
    monkeypatch.setattr(pipeline,'call_role',complete)
    first=pipeline.run_pipeline(tmp_path,RunRequest(pair_id='SYN-PAIR'))
    assert first.ai_analyses[0].status=='current'
    missing=first.checks[0].model_copy(deep=True)
    missing.check_id='MISSING-RELATION'
    first.checks.append(missing)
    seen=[]
    def retry(role,payload):
        seen.append((role,[c['check_id'] for c in payload['candidate_checks']]))
        assert all(c['check_id']=='MISSING-RELATION' for c in payload['candidate_checks'])
        item=opinion(run).model_dump()
        item['check_id']='MISSING-RELATION'
        result={'checks':[]} if role=='A2' else {'findings':[], 'analyses':[item]}
        return result, {'role':role,'performed':True,'status':'completed','model':'OFFLINE_TEST_DOUBLE'}
    monkeypatch.setattr(pipeline,'call_role',retry)
    second=pipeline.run_pipeline(tmp_path,RunRequest(pair_id='SYN-PAIR'),resume=first)
    assert [role for role,_ in seen]==['A2','A3']
    assert len(second.provider_calls)==2 and len(second.prior_provider_calls)==3
    assert len(second.ai_analyses)==2
    assert second.ai_analyses[0].inherited_from_run_id==first.run_id
    assert second.overview['model_opinions']['reused']==1


def test_malformed_model_item_does_not_erase_valid_neighbor(monkeypatch):
    from app import connectivity_pipeline as pipeline
    run=bridge_run()
    monkeypatch.setenv('DEEPSEEK_API_KEY','OFFLINE_TEST_ONLY')
    monkeypatch.setenv('AUDITTRACE_DEMO_USE_EXTERNAL_MODEL','true')
    class Response:
        def __enter__(self): return self
        def __exit__(self,*args): return False
        def read(self):
            arguments={'analyses':[opinion(run).model_dump(), {'summary':'未绑定核对项的坏条目'}], 'findings':[]}
            return json.dumps({'model':'OFFLINE_TEST_DOUBLE','choices':[{'message':{'tool_calls':[{'function':{
                'name':'submit_connectivity_relationships','arguments':json.dumps(arguments)}}]}}], 'usage':{}}).encode()
    monkeypatch.setattr(pipeline,'urlopen',lambda *args,**kwargs: Response())
    result,record=pipeline.call_role('A3',{'occurrences':[o.model_dump() for o in run.occurrences],
        'candidate_checks':[run.checks[0].model_dump()]})
    assert len(result['analyses'])==1 and record['validation_rejections']


def test_context_budget_preserves_members_and_reports_actual_coverage():
    run=bridge_run()
    rows=[o.model_dump() for o in run.occurrences]
    for index in range(120):
        extra=rows[0].copy()
        extra.update(occurrence_id=f'OPTIONAL-{index}',context='无关补充文字'*1200,row_header='补充信息')
        rows.append(extra)
    packed=pack_role_context({'occurrences':rows,'candidate_checks':[run.checks[0].model_dump()], 'calculated_results':[]})
    ids={o['occurrence_id'] for o in packed['occurrences']}
    assert {m.occurrence_id for m in run.checks[0].members}.issubset(ids)
    assert 'ADJ' in ids
    assert len(json.dumps(packed,ensure_ascii=False))<62500
    assert packed['actual_context_count']==len(ids)
    assert packed['context_packing']['omitted_occurrences']>0


def test_model_echo_of_server_fields_is_discarded_and_recomputed(monkeypatch):
    from app import connectivity_pipeline as pipeline
    run=bridge_run()
    monkeypatch.setenv('DEEPSEEK_API_KEY','OFFLINE_TEST_ONLY')
    monkeypatch.setenv('AUDITTRACE_DEMO_USE_EXTERNAL_MODEL','true')
    class Response:
        def __enter__(self): return self
        def __exit__(self,*args): return False
        def read(self):
            item=run.checks[0].model_dump()
            item['calculation']={'residual':'99999'}
            arguments={'checks':[item],'notes':[]}
            return json.dumps({'model':'OFFLINE_TEST_DOUBLE','choices':[{'message':{'tool_calls':[{'function':{
                'name':'submit_connectivity_relationships','arguments':json.dumps(arguments)}}]}}], 'usage':{}}).encode()
    monkeypatch.setattr(pipeline,'urlopen',lambda *args,**kwargs:Response())
    result,record=pipeline.call_role('A1',{'occurrences':[o.model_dump() for o in run.occurrences]})
    assert result['checks'] and 'calculation' not in result['checks'][0]
    assert {'hypothetical','scenario_adjustment','calculation'}.issubset(record['discarded_server_fields'])


def test_quota_gate_is_specific_stops_futile_roles_and_preserves_calculation(monkeypatch,tmp_path):
    from fastapi import HTTPException
    pipeline=offline_pipeline(monkeypatch,bridge_run())
    def deny():
        raise HTTPException(429,'本访客短时额度已用完，请稍后再试。')
    def never(*args,**kwargs):
        raise AssertionError('额度拒绝后不应请求供应商')
    monkeypatch.setattr(pipeline,'call_role',never)
    result=pipeline.run_pipeline(tmp_path,RunRequest(pair_id='SYN-PAIR'),before_model=deny)
    assert result.checks and not result.ai_analyses
    assert len(result.provider_calls)==1 and not result.provider_calls[0]['performed']
    assert result.provider_calls[0]['failure_code']=='MODEL_QUOTA_UNAVAILABLE'
    assert '访客' in result.provider_calls[0]['failure_message']
    assert result.agent_steps[-1]['status']=='skipped'


def test_old_prompt_opinion_stays_stale_after_numeric_refresh():
    run=bridge_run()
    integrate_analyses(run,[opinion(run)],{o.occurrence_id for o in run.occurrences},'A3','old-prompt')
    refresh_analysis_status(run,'current-prompt')
    assert run.ai_analyses[0].status=='needs_refresh'


def test_single_retry_is_analysis_only_and_retains_relationship_identity(monkeypatch,tmp_path):
    run=bridge_run()
    pipeline=offline_pipeline(monkeypatch,run)
    run.execution_mode='single_model'
    seen=[]
    def analyse(role,payload):
        assert payload['analysis_only'] is True
        seen.extend(c['check_id'] for c in payload['candidate_checks'])
        return {'checks':[],'analyses':[opinion(run).model_dump()]},{'role':role,'performed':True,'status':'completed','model':'OFFLINE_TEST_DOUBLE'}
    monkeypatch.setattr(pipeline,'call_role',analyse)
    result=pipeline.run_pipeline(tmp_path,RunRequest(pair_id='SYN-PAIR',mode='single_model'),resume=run)
    assert seen==[run.checks[0].check_id]
    assert {c.check_id for c in result.checks}=={c.check_id for c in run.checks}
    assert result.ai_analyses[0].status=='current'


def test_failed_usage_settlement_is_degraded_and_keeps_actual_analysis(monkeypatch,tmp_path):
    import time
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from app import connectivity_api
    from app.public_model import PublicModelQuotaError
    run=bridge_run()
    run.model_status='single_completed'
    doc=ReportDocument(document_id='SYN',case_id='SYN-PAIR',company_id='SYN',company='合成企业',report_type='AR',period='2025',source_url='https://example.invalid/synthetic')
    monkeypatch.setattr(connectivity_api,'load_registry',lambda root:[doc])
    monkeypatch.setattr(connectivity_api,'demo_task_supabase_enabled',lambda:False)
    def pipeline(*args,**kwargs):
        kwargs['before_model']()
        return run.model_copy(deep=True)
    monkeypatch.setattr(connectivity_api,'run_pipeline',pipeline)
    class Exhausted:
        def settle(self,*args,**kwargs):
            raise PublicModelQuotaError('本日输入预算已用完',code='PUBLIC_QUOTA_INPUT_TOKENS')
    app=FastAPI()
    app.include_router(connectivity_api.build_router(tmp_path,lambda request:'offline-lease',lambda:Exhausted()))
    with TestClient(app) as client:
        client.get('/api/connectivity/bootstrap')
        task_id=client.post('/api/connectivity/runs',json={'pair_id':'SYN-PAIR','mode':'single_model'}).json()['task_id']
        for _ in range(100):
            task=client.get('/api/connectivity/runs/'+task_id).json()
            if task['status'] not in {'queued','running'}:break
            time.sleep(.02)
        assert task['status']=='degraded'
        assert task['result']['checks']
        assert task['result']['model_status']=='single_completed'
        assert task['result']['versions']['quota_settlement']['code']=='PUBLIC_QUOTA_INPUT_TOKENS'
