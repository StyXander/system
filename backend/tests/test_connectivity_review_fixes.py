"""复核报告的业务反例；只验证机制，不伪装真实模型运行。"""
from copy import deepcopy
from io import BytesIO
from decimal import Decimal
import json
import csv
import pytest
from docx import Document
from app.connectivity_conditions import upgrade_conditions
from app.connectivity_engine import calculate
from app.connectivity_intelligence import integrate_analyses
from app.connectivity_overview import build_overview
from app.connectivity_delivery import export_run
from app.ai_opinion import build_program_hints,program_hint
from app.connectivity_schemas import DisclosureOccurrence,EvidenceRef,RunRequest
from test_connectivity_hybrid import bridge_run,opinion
from test_connectivity_hybrid_upgrade import offline_pipeline


def scope_run():
    run=bridge_run()
    for row in run.occurrences[:2]:row.label='环保投入';row.scope='集团'
    return run


def append_scope(run,proper=False):
    rows=[]
    found={o.occurrence_id:o for o in run.occurrences}
    for i,member in enumerate(found[m.occurrence_id] for m in run.checks[0].members):
        text='环保投入统计范围为集团，包含纳入环境数据统计的控股子公司。' if proper else ('上述酒类的统计口径均为商品酒。' if i==0 else '本报告披露范围与公司合并财务报表范围保持一致。')
        row=DisclosureOccurrence(occurrence_id=f'SCOPE-{i}',document_id=member.document_id,pdf_page=5,label='范围说明',context=text,quality='text_only')
        rows.append(row)
    run.occurrences.extend(rows)
    return [EvidenceRef(occurrence_id=r.occurrence_id,quote=r.context) for r in rows]


@pytest.mark.parametrize('member_scope',['','集团'])
def test_unrelated_product_scope_and_report_declaration_never_support_metric(member_scope):
    run=scope_run()
    for row in run.occurrences[:2]:row.scope=member_scope
    append_scope(run)
    check,upgraded=upgrade_conditions(run.checks[0],run,'2025')
    condition=next(c for c in check.conditions if c.name=='范围')
    assert condition.status=='unknown' and '范围' not in upgraded
    assert condition.evidence_kind=='report_level'


def test_even_model_supported_scope_must_match_current_metric_and_cited_quote():
    run=scope_run();refs=append_scope(run)
    condition=next(c for c in run.checks[0].conditions if c.name=='范围')
    condition.status='supported';condition.evidence=refs;condition.evidence_kind='ai_semantic'
    check,_=upgrade_conditions(run.checks[0],run,'2025')
    assert next(c for c in check.conditions if c.name=='范围').status=='unknown'


def test_metric_specific_scope_can_still_be_supported():
    run=scope_run();append_scope(run,proper=True)
    check,upgraded=upgrade_conditions(run.checks[0],run,'2025')
    assert '范围' in upgraded
    condition=next(c for c in check.conditions if c.name=='范围')
    assert condition.status=='supported' and condition.evidence_kind=='explicit'
    assert all('环保投入统计范围' in ref.quote for ref in condition.evidence)


def test_a3_retrieval_is_recalculated_before_model_and_stays_current(monkeypatch,tmp_path):
    run=scope_run();pipeline=offline_pipeline(monkeypatch,run)
    seen=[]
    def retrieve(root,current,checks,round_number):
        if round_number==2:
            append_scope(current,proper=True)
        return []
    monkeypatch.setattr(pipeline,'retrieve_targeted_pages',retrieve)
    def call(role,payload):
        seen.append(role)
        if role=='A3':
            scope=next(c for c in payload['candidate_checks'][0]['conditions'] if c['name']=='范围')
            assert scope['status']=='supported'
            assert payload['calculated_results'][0]['calculation']['residual'] is not None
            assert payload['calculated_results'][0]['hypothetical'] is None
            result={'findings':[],'analyses':[opinion(run).model_dump()]}
        else:result={'checks':[]}
        return result,{'role':role,'performed':True,'status':'completed','model':'OFFLINE_TEST_DOUBLE'}
    monkeypatch.setattr(pipeline,'call_role',call)
    result=pipeline.run_pipeline(tmp_path,RunRequest(pair_id='SYN-PAIR'))
    assert seen==['A1','A2','A3'] and result.model_status=='three_agents_completed'
    assert result.ai_analyses[0].status=='current'


@pytest.mark.parametrize('failure',['dimension','role'])
def test_hard_numeric_failure_cannot_be_consistent_or_strong(failure):
    run=bridge_run()
    if failure=='dimension':run.occurrences[1].unit='人'
    else:run.checks[0].members=run.checks[0].members[:1]
    run.checks[0]=calculate(run.checks[0],run.occurrences)
    assert run.checks[0].calculation_status=='not_executable'
    item=opinion(run);item.verdict='lean_consistent'
    ids={o.occurrence_id for o in run.occurrences}
    delta=integrate_analyses(run,[item],ids,'A3','offline')
    assert not delta['accepted_analysis_ids'] and delta['analysis_rejections'][0]['field']=='verdict'
    item.verdict='not_comparable';item.summary='当前成员无法执行数值比较，应先校正成员绑定再检验口径。'
    assert integrate_analyses(run,[item],ids,'A3','offline')['accepted_analysis_ids']
    assert run.ai_analyses[0].analysis_quality=='low'


def test_counts_separate_formal_trial_coverage_and_actual_program_hint():
    run=bridge_run();base=run.checks[0]
    formal=base.model_copy(deep=True);formal.check_id='FORMAL';formal.hypothetical=None
    formal.calculation={'residual':'0','left':'1','right':'1'};formal.calculation_status='consistent'
    formal.criteria_status='supported';formal.conclusion_level='formal_consistent'
    trial=base.model_copy(deep=True);trial.check_id='TRIAL';trial.calculation={'residual':None}
    trial.hypothetical={'residual':'-1','left':'1','right':'2','normalized_unit':'元'}
    # 明确构造有假设支持的试算，不能沿用base的依据不足级别冒充条件相容。
    trial.conclusion_level='assumption_compatible'
    text=base.model_copy(deep=True);text.check_id='TEXT';text.operator='narrative';text.hypothetical=None
    text.calculation={};text.calculation_status='not_applicable';text.conclusion_level='not_applicable'
    run.checks=[formal,trial,text];run.program_hints=build_program_hints(run)
    overview=build_overview(run)
    assert overview['formal_calculations']['items']==['FORMAL']
    assert overview['trial_calculations']['items']==['TRIAL']
    assert overview['numerical_records']['value']==2
    assert overview['program_records']['value']==3 and overview['program_hints']['value']==1
    csv_payload,_=export_run(run,'csv')
    count_rows={row[6]:int(row[7]) for row in csv.reader(csv_payload.decode('utf-8-sig').splitlines()) if row and row[0]=='overview_count'}
    assert count_rows['formal_calculations']==1 and count_rows['trial_calculations']==1
    assert count_rows['program_hints']==1 and count_rows['program_records']==3
    word,_=export_run(run,'docx');body='\n'.join(p.text for p in Document(BytesIO(word)).paragraphs)
    assert '正式复算1项' in body and '仅条件试算1项' in body and '程序提示3项' not in body


def test_relative_base_is_fixed_zero_is_unavailable_and_multiple_is_not_percentage_points():
    run=bridge_run();check=run.checks[0]
    check.hypothetical=None;check.calculation={'left':'100','right':'200','residual':'-100','inputs':[{'unit':'倍','dimension':'ratio'}]}
    diff=program_hint(check)['difference']
    assert Decimal(diff['relative_percent'])==100 and diff['relative_base']=='100'
    assert diff['percentage_points'] is None
    check.calculation.update(left='0.1',right='0.2',residual='-0.1',inputs=[{'unit':'%','dimension':'ratio'}])
    assert Decimal(program_hint(check)['difference']['percentage_points'])==Decimal('-10')
    check.calculation.update(left='0',right='0',residual='0')
    run.program_hints=build_program_hints(run)
    payload,_=export_run(run,'docx');body='\n'.join(p.text for p in Document(BytesIO(payload)).paragraphs)
    assert 'None%' not in body and '相对差异不适用' in body


def test_formal_recalculation_clears_old_hypothesis_and_bad_neighbor_does_not_abort():
    from app.connectivity_pipeline import recalculate_checks
    run=scope_run();append_scope(run,proper=True)
    bad=run.checks[0].model_copy(deep=True);bad.check_id='BAD';bad.members.append(bad.members[0])
    run.checks.append(bad)
    recalculate_checks(run,'2025','A3')
    assert len(run.checks)==1 and run.checks[0].hypothetical is None
    assert any(r['check_id']=='BAD' for r in run.candidate_rejections)


def test_coordinates_cannot_support_non_overlap():
    from test_connectivity_engine import relationship
    from app.connectivity_schemas import ConnectivityRun
    rows,check=relationship('RD-001',['100','60','40'],['total','component','component'])
    for row,label,box in zip(rows,['研发投入','费用化研发投入','资本化研发投入'],[[0,0,10,10],[0,20,10,30],[0,40,10,50]]):
        row.label=label;row.column_header='2025';row.bbox=box
    rows[0].context='研发投入包括费用化研发投入和资本化研发投入。'
    for condition in check.conditions:
        if condition.name in {'组成完整','互不重叠'}:condition.status='unknown';condition.evidence=[]
    run=ConnectivityRun(run_id='SYN',company='合成企业',period='2025',documents=[],occurrences=rows,checks=[check])
    fixed,_=upgrade_conditions(check,run,'2025')
    assert next(c for c in fixed.conditions if c.name=='互不重叠').status=='unknown'
    rows[0].context+='两者互不重叠。'
    fixed,_=upgrade_conditions(check,run,'2025')
    assert next(c for c in fixed.conditions if c.name=='互不重叠').status=='supported'


def test_current_module_exposes_no_fixed_verdict_entrypoint():
    from app import ai_opinion
    assert not hasattr(ai_opinion,'opinion_for_check') and not hasattr(ai_opinion,'merge_model_items')


def test_invalid_role_budget_falls_back_without_breaking_analysis(monkeypatch,tmp_path):
    run=bridge_run();pipeline=offline_pipeline(monkeypatch,run)
    monkeypatch.setenv('CONNECTIVITY_MAX_ROLE_BATCHES','bad-value')
    def call(role,payload):
        return ({'findings':[],'analyses':[opinion(run).model_dump()]} if role=='A3' else {'checks':[]}),{'role':role,'performed':True,'status':'completed','model':'OFFLINE_TEST_DOUBLE'}
    monkeypatch.setattr(pipeline,'call_role',call)
    result=pipeline.run_pipeline(tmp_path,RunRequest(pair_id='SYN-PAIR'))
    assert result.model_status=='three_agents_completed'
    assert any('默认4批' in item.get('reason','') for item in result.intelligence_trace)


def test_batch_diagnostic_is_not_saved_as_ai_summary():
    run=bridge_run()
    item=opinion(run)
    item.summary='右侧occurrence_id不在本批occurrences中'
    item.reasons=['右侧occurrence_id不在本批occurrences中','引文内容为比例行而非研发投入金额']
    item.next_actions=['核对右侧成员是否为研发投入金额行']
    delta=integrate_analyses(run,[item],{o.occurrence_id for o in run.occurrences},'A3','offline')
    assert delta['accepted_analysis_ids']
    saved=run.ai_analyses[0]
    assert 'occurrence_id' not in saved.summary
    assert all('occurrence_id' not in text for text in saved.reasons+saved.next_actions)
    assert saved.summary==saved.reasons[0]
    assert any(r.get('field')=='summary' for r in delta['analysis_rejections'])


def test_only_batch_diagnostics_is_rejected_with_reason():
    run=bridge_run()
    item=opinion(run)
    item.summary='右侧occurrence_id不在本批occurrences中'
    item.reasons=['右侧occurrence_id不在本批occurrences中']
    delta=integrate_analyses(run,[item],{o.occurrence_id for o in run.occurrences},'A3','offline')
    assert not delta['accepted_analysis_ids'] and not run.ai_analyses
    assert any('全部被移除' in r.get('reason','') for r in delta['analysis_rejections'])
