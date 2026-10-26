"""事项结果导出；直接读取保存版本，不再次调用模型。

所有格式读取同一份保存结果，不在导出时追加模型请求。
事项身份、路径、承诺、竞争解释与补证任务保持相同版本。
来源失效原因保留在引用旁，历史原文不冒充当前有效证据。
金额条件试算和模型文字判断不混入正式数值结论。
调用记录明确真实请求、失败状态与未知用量，不估算虚假费用。
表格单元格按原有导出转义规则处理，避免公式文本被执行。
"""
from __future__ import annotations

import csv
from io import BytesIO,StringIO
import json

from docx import Document
from docx.shared import Pt,RGBColor,Cm


def export_matter_run(run,format_name):
    from .connectivity_delivery import _cell,_add_heading,_paragraph,_table_rows
    from .delivery import _set_run_font
    notice=run.ai_generated_content_notice
    if format_name=="csv":
        stream=StringIO(newline="");writer=csv.writer(stream)
        writer.writerow(["记录类型","运行ID","版本","事项或证据ID","保存结果JSON"])
        writer.writerow(["metadata",run.run_id,run.result_revision,"",_cell(json.dumps({"notice":notice,"report_set":run.report_set.model_dump(mode="json"),"model_status":run.model_status,"overview":run.matter_overview},ensure_ascii=False))])
        for field in ("matters","matter_events","matter_commitments","matter_paths","matter_relations","matter_explanations",
                      "matter_assessments","matter_evidence","matter_tasks","matter_tool_receipts","matter_revisions",
                      "matter_coverage","matter_rejections","occurrences","provider_calls","prior_provider_calls"):
            for value in getattr(run,field):
                row=value.model_dump(mode="json") if hasattr(value,"model_dump") else value
                writer.writerow([field,run.run_id,run.result_revision,_cell(row.get("matter_id",row.get("evidence_id",""))),_cell(json.dumps(row,ensure_ascii=False))])
        return stream.getvalue().encode("utf-8-sig"),"text/csv; charset=utf-8"
    if format_name!="docx":raise ValueError("不支持的事项导出格式")
    doc=Document();section=doc.sections[0]
    section.top_margin=section.bottom_margin=Cm(2)
    normal=doc.styles["Normal"];normal.font.size=Pt(10.5)
    normal.paragraph_format.line_spacing=1.15
    _add_heading(doc,run.company+" 可持续事项财务影响调查",0)
    _paragraph(doc,notice)
    _paragraph(doc,f"运行 {run.run_id}  结果版本 {run.result_revision}  模式 {run.execution_mode}  模型状态 {run.model_status}")
    _paragraph(doc,f"资料期间 {run.period}  观察截至 {run.report_set.as_of_date}  资料集合 {run.report_set.report_set_id}")
    _paragraph(doc,"事项、证据、时间线和意见均来自该保存版本；支持状态描述披露对具体主张的支持程度。")
    _table_rows(doc,[(label,str(run.matter_overview.get(key,0))) for label,key in
        (("事项数","matter_total"),("已深查","investigated"),("待调查","pending"),("综合意见数","assessment_count"))])
    evidence={b.evidence_id:b for b in run.matter_evidence}
    for matter in run.matters:
        _add_heading(doc,matter.name,1)
        _paragraph(doc,f"事项 {matter.matter_id}  身份 {matter.identity_status}  调查 {matter.investigation_status}")
        _paragraph(doc,"身份依据："+matter.identity_basis)
        assessments=[a for a in run.matter_assessments if a.matter_id==matter.matter_id]
        for a in assessments:
            _add_heading(doc,"当前判断与未决问题",2)
            _paragraph(doc,a.summary)
            _paragraph(doc,f"关联支持 {a.relation_support}  金额可比 {a.amount_comparability}  人工采用 {a.review_status}")
            _paragraph(doc,"承诺进展："+a.commitment_progress+"；财务阶段："+a.financial_stage)
            for label,items in (("假设",a.assumptions),("分歧",a.disagreements),("待补证",a.unresolved)):
                for item in items:_paragraph(doc,label+"："+item)
            for action in a.next_actions:_paragraph(doc,f"下一步 {action.target_role}：{action.question}；资料：{action.requested_material}；用途：{action.purpose}")
        if not assessments:_paragraph(doc,"当前未形成综合复核意见，候选披露不能视为已完成调查。")
        _add_heading(doc,"财务影响路径",2)
        for path in run.matter_paths:
            if path.matter_id!=matter.matter_id:continue
            _paragraph(doc,f"{path.financial_item}  支持 {path.support}")
            for step in path.steps:_paragraph(doc,f"{step.label}：{step.explanation}；{step.origin}；{step.support}；引用 {', '.join(step.evidence_ids)}")
        for edge in run.matter_relations:
            if edge.matter_id==matter.matter_id:_paragraph(doc,f"信息关系 {edge.source_id} → {edge.target_id}：{edge.relation_kind}；{edge.evidence_status}；{edge.assertion_origin}；{edge.explanation}；引用 {', '.join(edge.evidence_ids)}")
        _add_heading(doc,"承诺与财务两线追踪",2)
        events=sorted([e for e in run.matter_events if e.matter_id==matter.matter_id],key=lambda e:(e.report_period,e.disclosure_date,e.event_id))
        for event in events:_paragraph(doc,f"披露期 {event.report_period}  发生期 {event.occurred_period or '未说明'}  {event.kind}：{event.label}；引用 {', '.join(event.evidence_ids)}")
        for commitment in run.matter_commitments:
            if commitment.matter_id==matter.matter_id:
                _paragraph(doc,f"承诺 {commitment.commitment_id}  基期 {commitment.baseline_year or '未明确'}  目标期 {commitment.target_year or '未明确'}  {commitment.status}；{commitment.label}")
        _add_heading(doc,"边界与竞争解释",2)
        for h in run.matter_explanations:
            if h.matter_id==matter.matter_id:
                _paragraph(doc,f"{h.status}：{h.explanation}；区分证据：{h.discriminating_evidence}；支持 {', '.join(h.supporting_evidence_ids)}；反证 {', '.join(h.counter_evidence_ids)}")
        _add_heading(doc,"事项原文",2)
        from .connectivity_matters import _all_refs
        relevant=set(matter.evidence_ids)
        for field in ("matter_paths","matter_events","matter_commitments","matter_explanations","matter_assessments"):
            for value in getattr(run,field):
                if value.matter_id==matter.matter_id:relevant.update(_all_refs(value.model_dump()))
        for eid in sorted(relevant):
            if eid not in evidence:continue
            b=evidence[eid]
            _paragraph(doc,f"{eid}  {b.document_id}  {b.report_period} {b.report_type}  PDF {b.pdf_page}  {b.extraction_method}")
            _paragraph(doc,b.quote)
            if b.invalid_reason:_paragraph(doc,"来源失效："+b.invalid_reason)
    _add_heading(doc,"覆盖与实际调用记录",1)
    for item in run.matter_coverage:_paragraph(doc,json.dumps(item,ensure_ascii=False))
    for task in run.matter_tasks:_paragraph(doc,f"{task.get('role')} {task.get('matter_id','')} {task.get('status')}：{task.get('stop_reason','')}")
    for call in run.prior_provider_calls:_paragraph(doc,"继承的历史调用："+json.dumps(call,ensure_ascii=False))
    for call in run.provider_calls:_paragraph(doc,"本运行调用："+json.dumps(call,ensure_ascii=False))
    _add_heading(doc,"登记原件",1)
    for source in run.documents:_paragraph(doc,f"{source.document_id}；{source.report_type} {source.effective_period}；发布 {source.disclosure_date}；SHA-256 {source.file_sha256}；{source.source_url}")
    for paragraph in doc.paragraphs:
        for r in paragraph.runs:_set_run_font(r)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for r in paragraph.runs:_set_run_font(r)
    out=BytesIO();doc.save(out)
    return out.getvalue(),"application/vnd.openxmlformats-officedocument.wordprocessingml.document"
