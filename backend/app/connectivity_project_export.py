"""Read-only Baosteel report downloads from the same verified data shown on screen."""
from __future__ import annotations

import csv
from datetime import datetime, timezone
from html import escape
from io import BytesIO, StringIO
import json
import re


FORMATS = {
    "json": "application/json; charset=utf-8",
    "csv": "text/csv; charset=utf-8",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "pdf": "application/pdf",
}
STATUS = {
    "completed": "已完成", "failed": "失败", "partial": "部分完成",
    "not_run": "未执行", "skipped": "跳过", "reused": "复用",
    "running": "运行中", "degraded": "部分完成",
    "five_agents_completed": "真实五 Agent 已完成",
    "five_agents_partial": "真实五 Agent 部分完成",
    "needs_evidence": "仍需证据", "blocked_without_scope": "口径未确认，暂不可比较",
}


def clean(value):
    # PDF extracts can contain XML-forbidden control characters (e.g. ESG bullet 0x07).
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", str(value if value is not None else "未取得"))


def label(value):
    return STATUS.get(value, clean(value))


def amount(value):
    return "未取得，未确认为零" if value is None else clean(value) + " 元"


def build_report(dossier, history):
    """No provider execution, no session-private history, no invented completion state."""
    return {
        "schema_version": "baosteel_special_report_v1",
        "title": "宝钢股份湛江钢铁近零碳项目专项核查报告",
        "exported_at": datetime.now(timezone.utc).isoformat(),
        "scope_notice": "本报告导出已登记原件核验结果及公开的历史模型回执；导出不触发新的模型调用。历史 AI 意见仅供参考，不构成审计结论。",
        "dossier": dossier,
        "featured_history": history,
    }


def report_blocks(report):
    """One content sequence for both readable formats; JSON retains the complete payload."""
    d = report["dossier"]
    blocks = []
    def add(kind, text):
        blocks.append((kind, clean(text)))
    add("title", report["title"])
    add("p", report["scope_notice"])
    add("p", "导出时间（UTC）：" + report["exported_at"])
    add("p", "资料索引时间：" + clean(d.get("generated_at")))
    add("h", "核查范围与核心问题")
    add("p", d.get("case_scope"))
    add("p", "项目：" + clean(d.get("project_label")))
    for i, q in enumerate(d.get("questions", []), 1):
        add("p", f"{i}. {q}")
    add("h", "原件核验与金额试算")
    add("p", f"已定位 {d.get('evidence_count', 0)} 处不同 PDF 页面证据。调查价值分级：{d.get('investigation_value', '未判定')}（不代表风险等级）。")
    add("p", d.get("mapping_note"))
    for reason in d.get("ranking_reasons", []):
        add("p", reason)
    r = d.get("reconciliation") or {}
    add("p", f"来源：{r.get('document_id', '未取得')} PDF 第 {r.get('source_page', '未取得')} 页；项目行：{r.get('source_row_label', '未取得')}")
    for key, name in (("opening", "期初余额"), ("increase", "本期增加"), ("to_fixed", "本期转固"), ("to_long_term", "转入长期待摊费用"), ("closing", "期末余额")):
        add("p", name + "：" + amount((r.get("values_yuan") or {}).get(key)))
    add("p", "试算公式：" + clean(r.get("formula")))
    add("p", "已知三项试算结余：" + amount(r.get("known_component_subtotal_yuan")))
    add("p", "计算状态：" + clean(r.get("status")) + "。" + clean(r.get("note")))
    add("h", "规则检查与待补证据")
    for item in d.get("rule_checks", []):
        add("sub", item.get("rule"))
        add("p", "状态：" + label(item.get("status")))
        add("p", item.get("question"))
    add("h", "官方原件证据清单")
    for i, e in enumerate(d.get("evidence", []), 1):
        add("sub", f"{i}. {e.get('document_id')}  PDF 第 {e.get('pdf_page')} 页")
        add("p", f"报告期间：{e.get('report_period')}；来源校验：{'已通过' if e.get('source_verified') else '未通过或未确认'}；证据编号：{e.get('evidence_id')}")
        add("p", "原文摘录：" + clean(e.get("quote")))
        add("p", "匹配词：" + "、".join(e.get("matched_terms") or []))
        add("p", "原件 SHA256：" + clean(e.get("source_sha256")))
    add("h", "后续核查程序")
    for i, item in enumerate(d.get("audit_procedures", []), 1):
        add("sub", f"{i}. {item.get('question')}")
        add("p", "所需材料：" + clean(item.get("material")))
        add("p", "目的：" + clean(item.get("purpose")))
    add("h", "真实五 Agent 历史运行记录")
    history = report["featured_history"]
    add("p", history.get("notice", "仅展示服务端白名单内的公开历史模型回执。"))
    if not history.get("items"):
        add("p", "当前未取得可展示的历史回执；不表示模型已经完成判断。")
    for run in history.get("items", []):
        add("sub", "任务 " + clean(run.get("task_id")))
        add("p", "任务创建时间：" + clean(run.get("created_at")))
        add("p", "模型状态：" + label(run.get("model_status")) + "；模型：" + "、".join(run.get("models") or []) + "；实际请求数：" + clean(run.get("actual_model_requests")))
        for role in run.get("role_results", []):
            add("p", f"{role.get('role')}：{label(role.get('status'))}；{role.get('failure_code') or role.get('identity_outcome') or role.get('stop_reason') or ''}")
        add("p", "历史 AI 综合意见（仅供参考）：" + clean(run.get("assessment_summary")))
        add("p", "财务阶段：" + clean(run.get("financial_stage")))
        add("p", f"财务路径 {run.get('financial_paths', 0)} 条；财务定位 {run.get('financial_occurrences', 0)} 处。")
        table = run.get("verified_table_detail")
        if table:
            add("p", f"该次运行独立核验原表：PDF 第 {table.get('source_page')} 页。" + clean(table.get("note")))
            add("p", "期初余额：" + amount((table.get("values_yuan") or {}).get("opening")) + "；本期增加：" + amount((table.get("values_yuan") or {}).get("increase")) + "；本期转固：" + amount((table.get("values_yuan") or {}).get("to_fixed")))
        for path in run.get("path_examples", []):
            add("p", "路径：" + clean(path.get("financial_item")) + "；证据等级：" + clean(path.get("support")))
            for step in path.get("steps", []):
                add("p", clean(step.get("label")) + "：" + clean(step.get("explanation")) + "；证据等级：" + clean(step.get("support")) + "；来源类型：" + clean(step.get("origin")))
        for item in run.get("counter_explanations", []):
            add("p", "替代解释：" + clean(item.get("explanation")) + "；状态：" + clean(item.get("status")))
            add("p", "判别证据：" + clean(item.get("required_evidence")))
        for item in run.get("next_actions", []):
            add("p", "待核问题：" + clean(item.get("question")))
            add("p", "所需材料：" + clean(item.get("requested_material")))
    add("h", "资料覆盖与使用边界")
    for c in d.get("coverage", []):
        add("p", f"{c.get('document_id')}：共 {c.get('total_pages')} 页；抽取失败页 {c.get('failed_pages', [])}；扫描页 {c.get('scanned_pages', [])}；OCR {'启用' if c.get('ocr_enabled') else '未启用'}。")
    add("p", (d.get("agents") or {}).get("notice"))
    for warning in d.get("source_warnings", []):
        add("p", warning)
    return blocks


def render_docx(blocks):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(.75)
    section.left_margin = section.right_margin = Inches(.85)
    for name in ("Normal", "Title", "Heading 1", "Heading 2"):
        style = document.styles[name]
        heading = name != "Normal"
        style.font.name = "黑体" if heading else "宋体"
        style.font.size = Pt(16 if heading else 12)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), style.font.name)
        properties = style._element.find(qn("w:pPr"))
        if properties is not None:
            for border in list(properties.findall(qn("w:pBdr"))):
                properties.remove(border)
    for kind, text in blocks:
        style = {"title": "Title", "h": "Heading 1", "sub": "Heading 2", "p": "Normal"}[kind]
        p = document.add_paragraph(style=style)
        p.paragraph_format.space_after = Pt(7)
        p.paragraph_format.line_spacing = 1.25
        p.paragraph_format.keep_with_next = kind != "p"
        for border in list(p._p.get_or_add_pPr().findall(qn("w:pBdr"))):
            p._p.pPr.remove(border)
        run = p.add_run(text)
        run.font.name = "宋体" if kind == "p" else "黑体"
        run.font.size = Pt(12 if kind == "p" else 16)
        run.font.color.rgb = RGBColor(0, 0, 0)
        run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), run.font.name)
    footer = section.footer.paragraphs[0]
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def render_pdf(blocks):
    import pymupdf
    tags = {"title": "h1", "h": "h2", "sub": "h3", "p": "p"}
    body = "".join(f"<{tags[kind]}>{escape(text).replace(chr(10), '<br>')}</{tags[kind]}>" for kind, text in blocks)
    # MuPDF supplies its bundled CJK fallback: no OS fonts or external services required.
    story = pymupdf.Story(html="<html><body>" + body + "</body></html>", user_css="""
        body { font-family: sans-serif; font-size: 11pt; color: #203b32; }
        h1 { font-size: 20pt; margin: 0 0 16pt; color: #173f32; }
        h2 { font-size: 15pt; margin: 18pt 0 10pt; color: #173f32; page-break-after: avoid; }
        h3 { font-size: 12pt; margin: 12pt 0 8pt; page-break-after: avoid; }
        p { margin: 0 0 8pt; line-height: 1.55; overflow-wrap: anywhere; }
    """)
    output = BytesIO()
    writer = pymupdf.DocumentWriter(output)
    page_box = pymupdf.paper_rect("a4")
    content_box = page_box + (48, 48, -48, -48)
    more = True
    while more:
        device = writer.begin_page(page_box)
        more, _ = story.place(content_box)
        story.draw(device)
        writer.end_page()
    writer.close()
    document = pymupdf.open(stream=output.getvalue(), filetype="pdf")
    document.set_metadata({"title": "宝钢专项核查报告", "author": "AuditTrace"})
    for i, page in enumerate(document):
        page.insert_text((page_box.width - 90, page_box.height - 25), f"{i + 1} / {len(document)}", fontsize=9)
    result = document.tobytes(garbage=4, deflate=True)
    document.close()
    return result


def export_report(format_name, dossier, history):
    if format_name not in FORMATS:
        raise ValueError("unsupported report format")
    report = build_report(dossier, history)
    if format_name == "json":
        return json.dumps(report, ensure_ascii=False, indent=2).encode("utf-8")
    if format_name == "csv":
        output = StringIO(newline="")
        writer = csv.writer(output)
        writer.writerow(["证据编号", "文档", "报告期间", "PDF页码", "原文摘录", "来源已校验", "SHA256"])
        for e in dossier.get("evidence", []):
            cells = [clean(e.get(k)) for k in ("evidence_id", "document_id", "report_period", "pdf_page", "quote", "source_verified", "source_sha256")]
            writer.writerow(["'" + v if v.lstrip().startswith(("=", "+", "-", "@")) else v for v in cells])
        return output.getvalue().encode("utf-8-sig")
    blocks = report_blocks(report)
    return render_docx(blocks) if format_name == "docx" else render_pdf(blocks)