"""同一保存结果的JSON、CSV与Word导出，保留口径和未决条件。

复用原delivery的字体、标题、表格方法，不另造排版框架。
公众摘要和专业明细共享同一版本事实，角色名称不替代运行证据。

导出合同的集成约定：
原始结果继续使用严格模型验证，新增字段提供空列表默认值。
历史记录缺少主张账本时保持未提供，不生成推测性证据。
五段阅读摘要与六级程序统计分别标明各自定义。
六级计数与浏览器共用显式结论字段及相同降级顺序。
有限支持仅说明结构或部分条件有依据，不能并入正式确认。
条件试算允许展示假设数值，但不填充正式残差字段。
披露位置以处计，解析覆盖以页计，不能合成为比例。
公共导出不包含个人观察正文，只保留已登记的数量。
表格中的来源文字仍按文本处理，避免电子表格执行公式。
规则条款尚未核验的标识必须保留到导出中。
所有标题与正文在保存前统一执行中文字体及字号要求。
"""
from __future__ import annotations

import csv
from io import BytesIO, StringIO
import json
import re

from docx import Document
from docx.shared import Pt, RGBColor, Cm

from .connectivity_schemas import ConnectivityRun, notice_for
from .delivery import _add_heading as _shared_heading, _table_rows as _shared_table, _set_run_font
from .connectivity_overview import build_overview
from .connectivity_facts import execution_provenance

STATUS_LABELS = {"supported": "依据齐备", "unknown": "依据待补", "conflicting": "存在冲突",
    "consistent": "所选关系计算一致", "precision_explained": "可由已知精度解释",
    "adjustment_explained": "有据加性调节后与0相容", "unexplained_difference": "未解释差异",
    "not_comparable": "暂不可比较", "not_executable": "暂不可执行", "not_applicable": "不适用数值计算",
    "unreviewed": "未人工复核", "demo_modified": "演示更正", "read_confirmed": "使用者已阅读"}
AI_LABELS = {"lean_consistent": "倾向一致", "potential_conflict": "存在潜在矛盾", "plausible_explanation": "存在合理解释路径",
             "insufficient_basis": "依据不足但可继续核查", "not_comparable": "当前不可直接比较"}


def export_fingerprint(run: ConnectivityRun) -> list[dict]:
    """同源校验用指纹：check_id、两类状态与正式残差必须与页面渲染一致。"""
    return [{"check_id": c.check_id, "criteria_status": c.criteria_status,
             "calculation_status": c.calculation_status, "residual": c.calculation.get("residual")}
            for c in run.checks]


def verify_export_same_source(run: ConnectivityRun, payload: bytes) -> list[str]:
    """比对导出 JSON 与内存结果的指纹；不一致即视为发布失败，返回差异说明。"""
    exported = json.loads(payload.decode("utf-8"))
    rows = exported.get("checks") or []
    left = export_fingerprint(run)
    if len(rows) != len(left):
        return [f"导出关系数{len(rows)}与结果{len(left)}不一致"]
    problems = []
    for field in ('program_hints', 'analysis_history'):
        if exported.get(field, []) != getattr(run, field):
            problems.append(f'{field}与保存结果不一致')
    for expected, actual in zip(left, rows):
        for key in ("check_id", "criteria_status", "calculation_status"):
            if expected[key] != actual.get(key):
                problems.append(f"{expected['check_id']}的{key}不一致")
        if (expected["residual"] or None) != ((actual.get("calculation") or {}).get("residual") or None):
            problems.append(f"{expected['check_id']}的正式残差与页面不一致")
    if exported.get("ai_analyses", []) != [a.model_dump(mode="json") for a in run.ai_analyses]:
        problems.append("导出AI研判与保存结果不一致")
    for expected, actual in zip(run.checks, rows):
        for field in ("calculation_plan", "scenario_adjustment", "narrative_test"):
            if expected.model_dump(mode="json").get(field) != actual.get(field):
                problems.append(f"{expected.check_id}的{field}与保存结果不一致")
    return problems


def _word_text(value) -> str:
    """PDF可能含XML非法控制字符；仅净化Word显示，不改JSON中的原始证据。"""
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]", "", str(value))


def _add_heading(document, text, level=1):
    return _shared_heading(document, _word_text(text), level)


def _table_rows(document, rows):
    return _shared_table(document, [(_word_text(key), _word_text(value)) for key, value in rows])


def _paragraph(document, text):
    """每段明确设正文宋体小四，避免继承系统默认字体。"""
    run = document.add_paragraph().add_run(_word_text(text))
    _set_run_font(run)
    run.font.size = Pt(12)


def _cell(value):
    text = str(value if value is not None else "")
    # 报告标签与模型文字不作为电子表格公式执行。
    return "'" + text if text.lstrip().startswith(("=", "+", "-", "@", "\t", "\r")) else text


def _assert_notice(payload: bytes, notice: str) -> None:
    """导出必须包含当次业务对应的声明常量，否则拒绝生成，不靠调用方自觉。"""
    if notice[:10] not in payload.decode("utf-8", "ignore"):
        raise ValueError("导出缺少当次业务对应的声明，已拒绝生成")


def export_run(run: ConnectivityRun, format_name: str, include_personal: bool = True) -> tuple[bytes, str]:
    # 只更新展示计数定义，不重写历史计算或模型意见。
    run = run.model_copy(update={'overview': build_overview(run)})
    notice = notice_for(run.task_kind)
    if format_name == "json":
        body = run.model_dump(mode="json")
        body["ai_generated_content_notice"] = notice
        if not include_personal:
            body["personal_observation_count"] = len(run.professional_notes)
            body["professional_notes"] = []
        payload = json.dumps(body, ensure_ascii=False, indent=2).encode("utf-8")
        _assert_notice(payload, notice)
        return payload, "application/json"
    sources = {o.occurrence_id: o for o in run.occurrences}
    if format_name == "csv":
        stream = StringIO(newline="")
        writer = csv.writer(stream)
        writer.writerow(["行类型", "运行", "版本", "核对项", "规则", "事项", "角色或条件", "原值或状态", "单位", "文件", "PDF顺序页", "原文或说明", "算式", "差额", "口径", "计算", "复核"])
        for key in ('formal_calculations', 'trial_calculations', 'program_records', 'program_hints', 'model_opinions', 'priority_checks'):
            entry = run.overview[key]
            writer.writerow(['overview_count', run.run_id, run.result_revision, '', '', '', key, entry['value'], '项', '', '',
                             entry.get('label', key), '', '', '', '', ''])
        for key, entry in run.overview["program_levels"].items():
            writer.writerow(["program_level_count", run.run_id, run.result_revision, "", "", "", key,
                             entry["value"], "项", "", "", "六级互斥分类", "", "", "", "", ""])
        for check in run.checks:
            base = [run.run_id, run.result_revision, check.check_id, check.rule_id, check.title]
            tail = [check.calculation.get("expression", ""), check.calculation.get("residual") or "", check.criteria_status, check.calculation_status, check.review_status]
            gate = check.calculation.get("gate") or {}
            note = "；".join(check.unresolved)
            if gate:
                note = f"{gate.get('conclusion', '暂不能直接比较')}：缺少{'、'.join(gate.get('missing') or [])}；{note}"
            writer.writerow([_cell(v) for v in ["relation", *base, "", "", "", "", "", note, *tail]])
            for member in check.members:
                o = sources[member.occurrence_id]
                writer.writerow([_cell(v) for v in ["member", *base, member.role, o.value, o.unit, o.document_id, o.pdf_page, o.context, *tail]])
            for condition in check.conditions:
                writer.writerow([_cell(v) for v in ["condition", *base, condition.name, condition.status, "", "", "", condition.evidence_kind + "；" + condition.reason + "；" + "；".join(ref.quote for ref in condition.evidence), *tail]])
            for index, position in enumerate(check.calculation.get("positions") or [], start=1):
                writer.writerow([_cell(v) for v in ["position", *base, f"第{index}处比较", position.get("residual", ""), "",
                    "", "", f"是否精度相容：{position.get('precision_compatible')}", *tail]])
            for item in (check.calculation.get("adjustment") or {}).get("applied", []):
                writer.writerow([_cell(v) for v in ["adjustment", *base, item["effect_id"], f"{item['sign']}{item['amount']}", item["unit"], "", "", item["label"] + "；" + item["note"], *tail]])
            for item in (check.calculation.get("adjustment") or {}).get("rejected", []):
                writer.writerow([_cell(v) for v in ["adjustment_rejected", *base, item.get("adjustment_id", ""), "已拒绝", "", "", "", item["reason"], *tail]])
            if check.closure:
                writer.writerow([_cell(v) for v in ["closure_probe", *base, "假设推演", "非正式结论", "", "", "", json.dumps(check.closure, ensure_ascii=False), *tail]])
            if check.hypothetical:
                writer.writerow([_cell(v) for v in ["hypothetical_only", *base, "假设试算", "非正式结论", check.hypothetical.get("normalized_unit", ""), "", "",
                    check.hypothetical.get("expression", "") + "；残差" + str(check.hypothetical.get("residual")), check.hypothetical.get("expression", ""), check.hypothetical.get("residual") or "", "假设", "不进入已衔接统计", ""]])
            if check.explanation:
                writer.writerow([_cell(v) for v in ["explanation", *base, "", "", "", "", "", check.explanation, *tail]])
            if check.related_checks:
                writer.writerow([_cell(v) for v in ["related_checks", *base, "", "", "", "", "", ";".join(check.related_checks), *tail]])
            opinion = next((item for item in run.ai_opinions if item.get("check_id") == check.check_id), None)
            if opinion:
                band = opinion.get("band") or {}
                writer.writerow([_cell(v) for v in ["ai_opinion", run.run_id, run.result_revision, check.check_id, check.rule_id, check.title,
                    band.get("label", ""), f"风险{opinion.get('score')}", "", "", "",
                    "；".join([str(opinion.get("headline", ""))] + [str(x) for x in opinion.get("rationale") or []]),
                    "", "", check.criteria_status, check.calculation_status, check.review_status]])
            for kind, value in [("calculation_plan", check.calculation_plan.model_dump() if check.calculation_plan else None),
                                ("scenario_adjustment", check.scenario_adjustment), ("narrative_test", check.narrative_test)]:
                if value:
                    writer.writerow([_cell(v) for v in [kind, *base, "", "", "", "", "", json.dumps(value, ensure_ascii=False), *tail]])
            for opinion in (a for a in run.ai_analyses if a.check_id == check.check_id):
                writer.writerow([_cell(v) for v in ["ai_analysis", *base, opinion.status, AI_LABELS[opinion.verdict], "", "", "",
                    json.dumps(opinion.model_dump(), ensure_ascii=False), *tail]])
            for hint in (h for h in run.program_hints if h.get('check_id') == check.check_id):
                writer.writerow([_cell(v) for v in ['program_hint', *base, '核查优先级', hint['priority'], '', '', '',
                    json.dumps(hint, ensure_ascii=False), *tail]])
            for hypothesis in check.hypotheses:
                writer.writerow([_cell(v) for v in ['explanation_hypothesis', *base, hypothesis.status, '', '', '', '',
                    hypothesis.model_dump_json(), *tail]])
        for exclusion in run.extraction_exclusions:
            writer.writerow([_cell(v) for v in ["excluded_occurrence", run.run_id, run.result_revision, "", "", "", "",
                exclusion.value, exclusion.unit, exclusion.document_id, exclusion.pdf_page, exclusion.label + "；" + exclusion.reason]])
        for record in run.version_records:
            writer.writerow([_cell(v) for v in ["version_record", run.run_id, run.result_revision, "", "", "", record.version_role,
                record.value or "未登记数值", record.unit, record.document_id, "", record.label + "；" + record.effective_relation]])
        for change in run.version_changes:
            writer.writerow([_cell(v) for v in ["version_change_only", run.run_id, run.result_revision, "", "", "", change.label,
                change.display_change, change.unit, change.from_version_id, "", change.note]])
        if include_personal:
            for note in run.professional_notes:
                writer.writerow([_cell(v) for v in ["personal_observation", run.run_id, run.result_revision, note.check_id,
                    "", "", note.actor, "个人观察，非系统结论", "", "", "", note.observation + "；理由：" + note.reason]])
        elif run.professional_notes:
            writer.writerow(["personal_observation_count", run.run_id, run.result_revision, "", "", "", "", f"存在{len(run.professional_notes)}条个人观察", "", "", "", "公共导出不含个人观察正文", "", "", "", "", ""])
        for review in run.reviews:
            writer.writerow([_cell(v) for v in ["review_event", run.run_id, run.result_revision, review.check_id,
                "", "", review.actor, review.action, "", "", "", json.dumps(review.model_dump(), ensure_ascii=False)]])
        for revision in run.analysis_history:
            writer.writerow([_cell(v) for v in ['ai_analysis_revision', run.run_id, run.result_revision,
                revision.get('check_id', ''), '', '', '', '', '', '', '', json.dumps(revision, ensure_ascii=False)]])
        for kind in ("claim_ledger", "rule_basis", "next_evidence"):
            for entry in getattr(run, kind):
                writer.writerow([_cell(v) for v in [kind, run.run_id, run.result_revision, entry.get("check_id", ""), "", "", "", "", "", "", "", json.dumps(entry, ensure_ascii=False)]])
        writer.writerow(["notice", notice])
        payload = stream.getvalue().encode("utf-8-sig")
        _assert_notice(payload, notice)
        return payload, "text/csv; charset=utf-8"
    if format_name != "docx":
        raise ValueError("不支持的报告格式")
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Cm(2)
    normal = document.styles["Normal"]
    normal.font.name = "宋体"
    normal.font.size = Pt(12)
    title = document.add_paragraph(style="Title").add_run("年度报告与可持续发展报告关联核对")
    _set_run_font(title, "黑体")
    title.font.size = Pt(16)
    title.font.color.rgb = RGBColor(0, 0, 0)
    _paragraph(document, f"{run.company} · {run.period} · {run.run_id} · 结果版本{run.result_revision}")
    _paragraph(document, notice)
    _paragraph(document, execution_provenance(run)["label"])
    _add_heading(document, "阅读摘要", 2)
    _paragraph(document, "可确认关系与其他状态采用同一六级分类，以下各类互斥，分母为全部关系：")
    for key, label in [("formal", "正式复算且数值相容"), ("conditional", "条件试算"), ("limited", "有限支持"),
                       ("insufficient", "依据不足·不可比"), ("conflict", "明确口径冲突或未解释差额"),
                       ("not_applicable", "不适用数值计算")]:
        _paragraph(document, f"{label}：{run.overview['program_levels'][key]['value']} / {len(run.checks)}项。")
    _paragraph(document, "条件试算依赖明确列出的假设；报告层声明不证明指标范围一致。正式一致为0不代表没有问题。")
    _paragraph(document, f"本次从登记报告中发现{len(run.occurrences)}处候选及上下文，形成{len(run.checks)}项关系记录、{len(run.findings)}项待查问题。数字相符不替代定义与范围依据。")
    overview = run.overview
    _paragraph(document, f'程序事实记录{len(run.program_hints)}条（含无差异登记）；正式复算{overview["formal_calculations"]["value"]}项；'
               f'仅条件试算{overview["trial_calculations"]["value"]}项；程序差异或限制提示{overview["program_hints"]["value"]}项。')
    _paragraph(document, f'当前模型意见覆盖{overview["model_opinions"]["value"]}/{len(run.checks)}项关系；中高优先级{overview["priority_checks"]["value"]}项。覆盖不等于异常，数量不相加为成果总数。')
    if run.versions.get('intelligence') and run.versions.get('rules') not in {'2026-10-18.2', '2026-10-19.1', '2026-10-20.3', '2026-10-21.1'}:
        _paragraph(document, '历史规则结果：原范围支持与意见未按本次缺陷修复重新验证。此处保留原计算，不应引用其正式支持数量作为当前成果。')
    overview = run.overview or {}
    if overview:
        _add_heading(document, "结果概览（各带分母，状态可重叠）", 3)
        for key in ("confirmable", "unconfirmable", "unexplained"):
            item = overview.get(key) or {}
            _paragraph(document, f"{item.get('label', key)}：{item.get('value')} / {item.get('denominator')}。")
        not_found = overview.get("not_found") or {}
        _paragraph(document, f"未归属披露：{not_found.get('unattributed', {}).get('value', 0)}处 / 已提取{len(run.occurrences)}处；解析失败：{not_found.get('failed_pages', {}).get('value', 0)}页 / 总计{not_found.get('total_pages', 0)}页。")
        _paragraph(document, overview.get("note", ""))
    if run.ai_opinions:
        _add_heading(document, "AI意见（程序刻度＋规则/模型解释，仅供参考）", 2)
        _paragraph(document, "AI意见不改变口径状态、正式差额与已衔接统计，不构成审计结论、鉴证结论或鉴证意见。")
        for opinion in run.ai_opinions:
            band = opinion.get("band") or {}
            _add_heading(document, f"{band.get('label', '未分级')}风险 {opinion.get('score', '')}｜{opinion.get('target', '')}", 3)
            _table_rows(document, [("结论类型", opinion.get("verdict", "")), ("置信度", opinion.get("confidence", "")),
                ("生成方式", "模型复核" if opinion.get("generated_by") == "llm" else "规则推理"),
                ("建议人工复核", "是" if opinion.get("needs_human") else "否")])
            _paragraph(document, "意见：" + str(opinion.get("headline", "")))
            for line in opinion.get("rationale") or []:
                _paragraph(document, "依据：" + str(line))
            for line in opinion.get("alternatives") or []:
                _paragraph(document, "另一种可能：" + str(line))
            for line in opinion.get("recommended_actions") or []:
                _paragraph(document, "建议下一步：" + str(line))
            if opinion.get("tags"):
                _paragraph(document, "标签：" + "、".join(str(tag.get("name", "")) for tag in opinion.get("tags") or []))
            _paragraph(document, str(opinion.get("disclaimer", "")))
    _table_rows(document, [("实际运行方式", run.execution_mode), ("模型状态", run.model_status), ("发现方式", run.discovery_mode),
        ("资料身份", run.dataset_role), ("耗时毫秒", run.duration_ms), ("父运行", run.parent_run_id or "无")])
    _add_heading(document, "本次文件与实际搜索覆盖", 2)
    for doc in run.documents:
        _table_rows(document, [("文件", doc.document_id), ("公司和期间", f"{doc.company} {doc.period}"),
            ("报告所属期间", doc.report_period or doc.period), ("指标所属期间", doc.metric_period or "未登记，按各披露行年度列归属"),
            ("来源类别", doc.source_kind), ("版本", doc.version), ("公开出处", doc.source_url), ("编制依据登记", doc.compilation_basis)])
    for coverage in run.coverage:
        _paragraph(document, f"{coverage.document_id}：共{coverage.total_pages}页，解析{coverage.parsed_pages}页，搜索{len(coverage.searched_pages)}页；失败页：{json.dumps(coverage.failed_pages, ensure_ascii=False)}。仅说明实际处理范围。")
    seen_member_context = set()
    _paragraph(document, '本Word为同源阅读报告，原文上下文按首次引用节选；完整计算、执行明细及原文上下文见同一运行的JSON、CSV和原页。')
    for check in run.checks:
        _add_heading(document, check.title, 2)
        _table_rows(document, [("核对项及规则", f"{check.check_id} / {check.rule_id}"),
            ("口径状态", STATUS_LABELS.get(check.criteria_status, check.criteria_status)),
            ("计算状态", STATUS_LABELS.get(check.calculation_status, check.calculation_status)),
            ("复核状态", STATUS_LABELS.get(check.review_status, check.review_status))])
        for opinion in (a for a in run.ai_analyses if a.check_id == check.check_id):
            _add_heading(document, "AI分析意见：" + AI_LABELS[opinion.verdict], 3)
            _paragraph(document, "仅供核查参考；" + ("输入已变化，需重新分析" if opinion.status == "needs_refresh" else "与本次保存结果绑定"))
            _paragraph(document, "使用者处理：" + {"unreviewed": "未复核", "accepted": "已采纳为核查参考", "rejected": "已驳回"}[opinion.review_status])
            _paragraph(document, opinion.summary)
            _paragraph(document, f'生成角色：{opinion.producer_role}；实际模型：{opinion.producer_model or "历史记录未登记"}；' +
                       (f'复用原运行：{opinion.inherited_from_run_id}' if opinion.inherited_from_run_id else '本次模型生成'))
            if opinion.priority:
                _paragraph(document, 'AI核查优先级：' + {'high': '高', 'medium': '中', 'low': '低'}[opinion.priority] + '；' + opinion.priority_reason)
            for title, values in [("判断依据", opinion.reasons), ("假设", opinion.assumptions),
                                  ("其他解释", opinion.alternative_explanations), ("下一步", opinion.next_actions),
                                  ("分析质量说明", opinion.quality_reasons)]:
                if values:
                    _paragraph(document, title + "：" + "；".join(values))
            for kind, refs in [("支持原文", opinion.supporting_evidence), ("相反证据", opinion.counter_evidence)]:
                for ref in refs:
                    row = sources.get(ref.occurrence_id)
                    _paragraph(document, f"{kind} {ref.occurrence_id} / PDF {row.pdf_page if row else '?'}：{ref.quote}")
        if check.calculation_plan:
            _paragraph(document, "计算方案：" + check.calculation_plan.reason + "；假设：" + "；".join(check.calculation_plan.assumptions))
        for hint in (h for h in run.program_hints if h.get('check_id') == check.check_id):
            diff = hint['difference']
            _paragraph(document, '核查优先级：' + {'high': '高', 'medium': '中', 'low': '低'}[hint['priority']] + '；' + hint['priority_reason'])
            if diff['absolute'] is not None:
                basis_label = '左侧基准' if diff.get('relative_basis') == 'left' else '历史原算法基准'
                relative = f'相对差异{diff["relative_percent"]}%（{basis_label}）' if diff['relative_percent'] is not None else '基准为零或缺失，相对差异不适用'
                _paragraph(document, ('条件试算' if diff['conditional'] else '正式复算') + f'绝对差额{diff["absolute"]}{diff["unit"]}；{relative}。')
            _paragraph(document, '已支持：' + '、'.join(hint['coverage']['supported']) + '；待补：' + '、'.join(hint['coverage']['unknown']) +
                       '；存在冲突：' + '、'.join(hint['coverage']['conflicting']))
        for hypothesis in check.hypotheses:
            _paragraph(document, '竞争解释：' + hypothesis.explanation + '；状态：' +
                       {'supported': '有证据支持', 'plausible': '可能解释', 'open': '待验证', 'refuted': '已有反证'}[hypothesis.status])
            _paragraph(document, '依据：' + '；'.join(hypothesis.reasons) + '；假设：' + '；'.join(hypothesis.assumptions))
            for ref in hypothesis.supporting_evidence + hypothesis.counter_evidence:
                row = sources.get(ref.occurrence_id)
                _paragraph(document, f'原文 {ref.occurrence_id} / PDF {row.pdf_page if row else "?"}：{ref.quote}')
        if check.scenario_adjustment:
            _add_heading(document, "条件调节试算（不覆盖正式差额）", 3)
            bridge = check.scenario_adjustment
            _paragraph(document, f'原始候选差额：{bridge["original_residual"]}{bridge["unit"]}')
            for entry in bridge['applied']:
                _paragraph(document, f'{entry["label"]}：有符号贡献{entry["amount"]}{bridge["unit"]}；{entry["reason"]}')
                for ref in entry['evidence']:
                    _paragraph(document, f'依据 {ref["occurrence_id"]}：{ref["quote"]}')
            _paragraph(document, f'剩余候选差额：{bridge["remaining_residual"]}{bridge["unit"]}；假设：' + '；'.join(bridge['assumptions']))
        if check.narrative_test:
            _paragraph(document, "文字与数值方向核对：" + json.dumps(check.narrative_test, ensure_ascii=False))
        for condition in check.conditions:
            kind = {'report_level': '仅报告总体声明', 'explicit': '本指标原文明示', 'ai_semantic': 'AI语义解释', 'table_structure': '表格结构', 'derived': '程序派生'}.get(condition.evidence_kind, '原记录未分类')
            _paragraph(document, f"{condition.name}：{STATUS_LABELS.get(condition.status, condition.status)}；依据类型：{kind}；{condition.reason}")
            for ref in condition.evidence:
                _paragraph(document, f"依据 {ref.occurrence_id}：{ref.quote}")
        for member in check.members:
            row = sources[member.occurrence_id]
            _table_rows(document, [("成员与角色", f"{row.occurrence_id} / {member.role}"), ("原标签", row.label),
                ("原值与单位", f"{row.value if row.value is not None else '未提取数值'} {row.unit}"),
                ("位置", f"{row.document_id} PDF顺序页{row.pdf_page}；列头{row.column_header}；年度{row.year or '未归属'}（{row.year_status}）")])
            if row.occurrence_id not in seen_member_context:
                seen_member_context.add(row.occurrence_id)
                start = max(0, row.context.find(row.label) - 80)
                excerpt = row.context[start:start+520]
                _paragraph(document, ('原文上下文节选（空白收拢）：' if len(excerpt) < len(row.context) else '原文上下文（空白收拢）：') + re.sub(r'\s+', ' ', excerpt))
            else:
                _paragraph(document, '原文上下文已在本报告前文同ID成员处列示。')
        gate = check.calculation.get("gate") or {}
        if gate:
            _paragraph(document, f"正式结论：{gate.get('conclusion', '暂不能直接比较')}；缺少：{'、'.join(gate.get('missing') or [])}；{gate.get('needed', '')}")
        facts = {k:v for k,v in check.calculation.items() if k in
                 {'expression','left','right','residual','normalized_unit','residual_interval','precision_compatible','denominator_check'}}
        if facts:
            names = {'expression':'算式','left':'左侧','right':'右侧','residual':'差额','normalized_unit':'归一化单位',
                     'residual_interval':'差额区间','precision_compatible':'是否可由披露精度解释','denominator_check':'分母检查'}
            for key,value in facts.items():
                _paragraph(document, names[key] + '：' + (json.dumps(value,ensure_ascii=False) if isinstance(value,dict) else str(value)))
        if check.closure:
            _add_heading(document, "补证闭合推演（假设分析）", 3)
            probe = check.closure
            ceiling = probe.get('ceiling', {})
            _paragraph(document, f"{probe.get('tier_label', '')}：{probe.get('tier_reason', '')}")
            _paragraph(document, f"假定前提补齐后的上限：{ceiling.get('label', '')}；差额{ceiling.get('residual')} {ceiling.get('normalized_unit', '')}。不作为正式结论。")
            for demand in probe.get('demands', []):
                _paragraph(document, f"待补{demand.get('condition', '')}：{demand.get('procedure', '')}")
        if check.hypothetical:
            _add_heading(document, "假设试算（非正式结论，不进入已衔接统计）", 3)
            _paragraph(document, "假设：" + "；".join(check.hypothetical.get("assumptions") or []))
            _paragraph(document, f"试算算式：{check.hypothetical.get('expression', '')}；试算差额：{check.hypothetical.get('residual')} {check.hypothetical.get('normalized_unit', '')}")
        adjustment = check.calculation.get("adjustment") or {}
        if adjustment:
            _add_heading(document, "差异解释的加性调节与残差", 3)
            for item in adjustment.get("applied", []):
                _paragraph(document, f"调节 {item['adjustment_id']}（效果{item['effect_id']}）：{item['sign']}{item['amount']} {item['unit']}；{item['label']}。{item['note']}")
            for item in adjustment.get("rejected", []):
                _paragraph(document, f"已拒绝的调节 {item.get('adjustment_id')}：{item['reason']}")
            if adjustment.get("applied"):
                interval = adjustment.get("remaining_interval") or {}
                _paragraph(document, f"剩余残差：{adjustment.get('remaining_residual')}；区间[{interval.get('low')}, {interval.get('high')}]；是否含0：{interval.get('contains_zero')}。{adjustment.get('note', '')}")
        _paragraph(document, "解释：" + (check.explanation or "本次尚未找到充分解释"))
        _paragraph(document, "关联的解释或分组记录：" + ("；".join(check.related_checks) or "暂无"))
        _paragraph(document, "未决条件：" + ("；".join(check.unresolved) or "当前选定关系无额外未决条件"))
    if run.extraction_exclusions:
        _add_heading(document, '附录：已提取但未计入的位置', 2)
        _paragraph(document, f'共{len(run.extraction_exclusions)}处未进入候选，以下按文件页码归并列示；这些位置不作为没有披露的证据。')
        exclusions_by_page = {}
        for exclusion in run.extraction_exclusions:
            exclusions_by_page.setdefault((exclusion.document_id, exclusion.pdf_page), []).append(exclusion)
        for (document_id, page), exclusions in exclusions_by_page.items():
            _paragraph(document, f'{document_id} PDF顺序页{page}：' + '；'.join(
                f'{item.label} {item.value} {item.unit}（{item.reason}）' for item in exclusions))
    if run.version_records or run.version_changes:
        _add_heading(document, "企业披露版本与版本变化（不进入正式跨报告差额）", 2)
        for record in run.version_records:
            _paragraph(document, f"{record.document_id} · {record.version_role} · {record.disclosure_date or '披露日期未登记'} · 指标期间{record.metric_period or '未登记'} · {record.effective_relation} · 来源{record.source_kind}")
        for change in run.version_changes:
            _paragraph(document, f"{change.label}：{change.from_version_id} → {change.to_version_id}，显示变化{change.display_change} {change.unit}。{change.note}")
    _add_heading(document, "具体问题与后续资料用途", 2)
    for finding in run.findings:
        _table_rows(document, [("关联核对项", finding.check_id), ("已知事实", finding.known_facts), ("待查问题", finding.question),
            ("可进一步了解的说明或资料", finding.requested_material), ("用途", finding.purpose)])
    _add_heading(document, "使用者个人观察", 2)
    if not include_personal:
        _paragraph(document, f"公共导出不含个人观察正文；本次记录存在{len(run.professional_notes)}条个人观察，仅本人视图可展开。")
    for note in (run.professional_notes if include_personal else []):
        _paragraph(document, f"{note.check_id} / {note.actor} / {note.created_at}：{note.observation}；理由：{note.reason}。本项为个人观察，不改变系统计算结论。")
    _add_heading(document, "角色、调用及修订记录", 2)
    for step in run.agent_steps:
        _paragraph(document, f'{step.get("role")}：{step.get("status")}；新增关系{len(step.get("new_check_ids",[]))}项，'
                   f'更新关系{len(step.get("updated_check_ids",[]))}项，新增证据{len(step.get("added_evidence_ids",[]))}处，'
                   f'研判{len(set(step.get("accepted_analysis_ids",[])))}项，失败批次{len(step.get("failed_batches",[]))}。')
        for rejection in step.get('analysis_rejections',[]):
            _paragraph(document, f'意见局部未采纳 {rejection.get("check_id", "")} / {rejection.get("field", "")}：{rejection.get("reason", "")}')
    collaboration = overview.get('collaboration', {})
    _add_heading(document, "自主补证与反证回执", 2)
    _paragraph(document, collaboration.get('notice', '未记录自主工具执行'))
    for event in collaboration.get('events', []):
        _paragraph(document, f"{event.get('producer_role', event.get('role', '程序'))} / {event.get('kind')} / {event.get('check_id', '')}："
                   f"{event.get('reason', event.get('stop_reason', event.get('note', '')))}；"
                   f"工具{event.get('tool', '—')}；新增原文{len(event.get('new_evidence_ids', []))}处；条件升级关系{event.get('upgraded', '—')}。")
    for call in run.provider_calls:
        _paragraph(document, f'{call.get("role")} / {call.get("model", "未发出模型请求")} / {call.get("status")}；'
                   f'输入{call.get("input_tokens", "未知")}、输出{call.get("output_tokens", "未知")} tokens；'
                   f'耗时{call.get("duration_ms", "未知")}毫秒；{call.get("failure_message", call.get("failure_code", ""))}')
    for review in run.reviews:
        _paragraph(document, f'{review.created_at} / {review.actor} / {review.check_id}：{review.action}；{review.reason}')
    for revision in run.analysis_history:
        old = revision['previous']
        _paragraph(document, f'意见修订 {revision["check_id"]}：原意见“{old["summary"]}”；修订原因：{revision["reason"]}')
    for kind, label in [("claim_ledger", "AI主张证据账本"), ("rule_basis", "规则依据与核验状态"), ("next_evidence", "具体补证动作")]:
        if getattr(run, kind):
            _add_heading(document, label, 2)
            for entry in getattr(run, kind):
                _paragraph(document, json.dumps(entry, ensure_ascii=False))
    # 所有层级标题和表格也执行同一字号要求，不依赖Word默认样式。
    def format_paragraphs(paragraphs):
        for paragraph in paragraphs:
            style_id = paragraph._p.style or ''
            heading = style_id.startswith('Heading') or style_id == 'Title'
            for text_run in paragraph.runs:
                _set_run_font(text_run, '黑体' if heading else '宋体')
                text_run.font.size = Pt(16 if heading else 12)
    format_paragraphs(document.paragraphs)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                format_paragraphs(cell.paragraphs)
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
